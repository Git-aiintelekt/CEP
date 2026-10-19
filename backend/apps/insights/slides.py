"""PowerPoint exports (python-pptx): a report or a dashboard as a 16:9 deck that stays editable. A title slide, then one slide per
tile (or for the report's view): bars, lines, areas, stacks, pies, funnels and waterfalls as native PowerPoint charts with
their data, a scatter or a bubble as an XY chart, a figure, a KPI card or a gauge as its number against its target, a summary
or a text tile as words, and every other view as a table of its first rows.
"""
import io

from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .pdf import _number, cell_text, chart_data

WIDE, TALL = Inches(13.333), Inches(7.5)
INK = RGBColor(0x17, 0x18, 0x1C)
INK_700 = RGBColor(0x6B, 0x6F, 0x7B)
LINE = RGBColor(0xCF, 0xD2, 0xDC)
PALETTE = [RGBColor.from_string(c) for c in ("3863DB", "A8488F", "418EAE", "9EA2AD", "B08A6E", "7570D4")]
TABLE_ROWS, TABLE_COLUMNS = 12, 8
CHART_KINDS = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA, "stacked": XL_CHART_TYPE.COLUMN_STACKED,
               "stacked100": XL_CHART_TYPE.COLUMN_STACKED_100, "combo": XL_CHART_TYPE.COLUMN_CLUSTERED, "pie": XL_CHART_TYPE.PIE,
               "funnel": XL_CHART_TYPE.BAR_CLUSTERED, "waterfall": XL_CHART_TYPE.COLUMN_CLUSTERED}
ONE_SERIES = ("pie", "funnel", "waterfall")


def _text(slide, words: str, left, top, width, height, *, size: int = 14, bold: bool = False, color=INK, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for i, part in enumerate(str(words).split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = part
        paragraph.font.size, paragraph.font.bold, paragraph.font.color.rgb = Pt(size), bold, color
        if align is not None:
            paragraph.alignment = align
    return box


def _slide(deck, title: str, caption: str | None = None):
    slide = deck.slides.add_slide(deck.slide_layouts[6])                  # blank
    _text(slide, title, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.7), size=24, bold=True)
    if caption:
        _text(slide, caption, Inches(0.6), Inches(1.0), Inches(12.1), Inches(0.4), size=12, color=INK_700)
    return slide


BODY = (Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.6))


def _table(slide, columns: list[dict], rows: list[list], note: str | None = None):
    shown = columns[:TABLE_COLUMNS]
    body = rows[:TABLE_ROWS]
    left, top, width, height = BODY
    table = slide.shapes.add_table(len(body) + 1, len(shown), left, top, width, Emu(min(int(height), int(Inches(0.4)) * (len(body) + 1)))).table
    for c, column in enumerate(shown):
        cell = table.cell(0, c)
        cell.text = column["label"]
        cell.text_frame.paragraphs[0].font.size, cell.text_frame.paragraphs[0].font.bold = Pt(11), True
    for r, row in enumerate(body, start=1):
        for c, column in enumerate(shown):
            cell = table.cell(r, c)
            cell.text = _number(column, row[c]) if column.get("aggregate") else cell_text(column, row[c])
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.font.size = Pt(11)
            if column.get("aggregate"):
                paragraph.alignment = PP_ALIGN.RIGHT
    extra = []
    if len(rows) > TABLE_ROWS:
        extra.append(f"The first {TABLE_ROWS} of {len(rows):,} rows.")
    if len(columns) > TABLE_COLUMNS:
        extra.append(f"The first {TABLE_COLUMNS} of {len(columns)} columns.")
    if note:
        extra.insert(0, note)
    if extra:
        _text(slide, " ".join(extra), Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.35), size=10, color=INK_700)


def _figure(slide, label: str, value: str, lines: list[str]):
    _text(slide, label, Inches(0.6), Inches(2.2), Inches(12.1), Inches(0.5), size=16, color=INK_700, align=PP_ALIGN.CENTER)
    _text(slide, value, Inches(0.6), Inches(2.8), Inches(12.1), Inches(1.4), size=60, bold=True, align=PP_ALIGN.CENTER)
    if lines:
        _text(slide, "\n".join(lines), Inches(0.6), Inches(4.4), Inches(12.1), Inches(1.0), size=16, color=INK_700, align=PP_ALIGN.CENTER)


def _grand(columns: list[dict], rows: list[list], totals):
    entry = next((t for t in totals or [] if not t.get("keys")), None)
    if entry is not None and entry.get("values"):
        return entry["values"][0]
    measures = [i for i, c in enumerate(columns) if c.get("aggregate")]
    return rows[0][measures[0]] if len(rows) == 1 and measures else None


def _chart(slide, columns: list[dict], rows: list[list], view: dict) -> bool:
    kind = view.get("chart") or "bar"
    data = chart_data(columns, rows, None if kind in ONE_SERIES + ("combo",) else view.get("across"))
    if data is None or not data["categories"]:
        return False
    chart_data_ = CategoryChartData()
    chart_data_.categories = [label for _, label in data["categories"]]
    for label, values in (data["series"][:1] if kind in ONE_SERIES else data["series"]):
        chart_data_.add_series(label, values)
    left, top, width, height = BODY
    chart = slide.shapes.add_chart(CHART_KINDS[kind], left, top, width, height, chart_data_).chart
    chart.has_legend = kind == "pie" or len(data["series"]) > 1
    if chart.has_legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        chart.legend.font.size = Pt(11)
    if kind == "pie":
        for i, point in enumerate(chart.plots[0].series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = PALETTE[i % len(PALETTE)]
    else:
        for i, series in enumerate(chart.plots[0].series):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = PALETTE[i % len(PALETTE)]
            if kind == "line":
                series.format.line.color.rgb = PALETTE[i % len(PALETTE)]
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = LINE
        chart.value_axis.tick_labels.font.size = Pt(11)
        chart.category_axis.tick_labels.font.size = Pt(11)
        if kind == "funnel":
            chart.category_axis.reverse_order = True
    return True


def _xy(slide, columns: list[dict], rows: list[list], bubble: bool) -> bool:
    dims = [i for i, c in enumerate(columns) if not c.get("aggregate")]
    meas = [i for i, c in enumerate(columns) if c.get("aggregate")]
    if not dims or len(meas) < (3 if bubble else 2):
        return False
    data = BubbleChartData() if bubble else XyChartData()
    for row in rows[:200]:
        series = data.add_series(cell_text(columns[dims[0]], row[dims[0]]))
        x, y = float(row[meas[0]] or 0), float(row[meas[1]] or 0)
        if bubble:
            series.add_data_point(x, y, float(row[meas[2]] or 0))
        else:
            series.add_data_point(x, y)
    left, top, width, height = BODY
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BUBBLE if bubble else XL_CHART_TYPE.XY_SCATTER, left, top, width, height, data).chart
    chart.has_legend = len(rows) <= 12
    for i, series in enumerate(chart.plots[0].series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = PALETTE[i % len(PALETTE)]
    return True


def _view_slide(slide, columns: list[dict], rows: list[list], view: dict | None, totals):
    view = view or {"type": "grid"}
    kind = view.get("type", "grid")
    chart = view.get("chart") or "bar"
    if not rows:
        _text(slide, "No rows match this report.", *BODY, size=14, color=INK_700)
        return
    if kind == "chart" and chart in ("kpi", "gauge"):
        measure = next((c for c in columns if c.get("aggregate")), None)
        value = _grand(columns, rows, totals)
        lines = []
        if view.get("target") is not None and value is not None:
            lines.append(f"Target {_number(measure, view['target'])} · {float(value) / float(view['target']) * 100:.1f}% of target" if view["target"] else f"Target {_number(measure, view['target'])}")
        _figure(slide, measure["label"] if measure else "", _number(measure, value) if measure else "—", lines)
        return
    if kind == "chart" and chart in CHART_KINDS and _chart(slide, columns, rows, view):
        if chart == "combo":
            _text(slide, "A combo chart is shown with every value as bars.", Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.35), size=10, color=INK_700)
        return
    if kind == "chart" and chart in ("scatter", "bubble") and _xy(slide, columns, rows, chart == "bubble"):
        return
    note = {"treemap": "A treemap", "map": "A map", "influencers": "Key Influencers", "decomposition": "A Decomposition Tree"}.get(chart if kind == "chart" else kind)
    _table(slide, columns, rows, f"{note} is shown here as its rows." if note else None)


def _deck():
    deck = Presentation()
    deck.slide_width, deck.slide_height = WIDE, TALL
    return deck


def _cover(deck, name: str, subtitle: str, workspace: str, when: str):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _text(slide, name, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.2), size=40, bold=True)
    _text(slide, subtitle, Inches(0.8), Inches(3.8), Inches(11.7), Inches(0.6), size=18, color=INK_700)
    _text(slide, f"{workspace} · {when}", Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.4), size=12, color=INK_700)


def _bytes(deck) -> bytes:
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def report(name: str, columns: list[dict], rows: list[list], view: dict | None, totals=None, *, workspace: str, when: str) -> bytes:
    deck = _deck()
    _cover(deck, name, f"{len(rows):,} row{'s' if len(rows) != 1 else ''}", workspace, when)
    slide = _slide(deck, name)
    _view_slide(slide, columns, rows, view, totals)
    if (view or {}).get("type", "grid") != "grid" and rows:                        # the rows too, after a chart or a pivot
        _table(_slide(deck, f"{name}: the rows"), columns, rows)
    return _bytes(deck)


def dashboard(name: str, description: str | None, period_words: str, tiles: list[dict], *, workspace: str, when: str) -> bytes:
    """tiles: as for the PDF ({"title", "kind": "figure" | "view" | "summary" | "text" | "missing" | "refused", and the payload})"""
    deck = _deck()
    _cover(deck, name, f"{description + ' · ' if description else ''}{period_words}", workspace, when)
    for tile in tiles:
        kind = tile.get("kind")
        if kind == "figure":
            slide = _slide(deck, tile["title"])
            value, previous = tile.get("value"), tile.get("previous")
            lines = []
            if tile.get("compared") and previous not in (None, 0) and value is not None:
                change = (float(value) - float(previous)) / abs(float(previous))
                lines.append(f"{'Up' if change >= 0 else 'Down'} {abs(change) * 100:.1f}% on the previous period")
            _figure(slide, tile["column"]["label"], _number(tile["column"], value), lines)
        elif kind == "view":
            result = tile["result"]
            slide = _slide(deck, tile["title"], f"{result['row_count']:,} row{'s' if result['row_count'] != 1 else ''}")
            _view_slide(slide, result["columns"], result["rows"], tile["view"], result.get("totals"))
        elif kind in ("summary", "text"):
            slide = _slide(deck, tile["title"], "Written by AI from the rows" if kind == "summary" and tile.get("summary") else None)
            words = tile.get("text") if kind == "text" else "\n".join([tile.get("summary") or tile.get("detail") or "No summary yet.", *(f"• {p}" for p in tile.get("points") or [])])
            _text(slide, words, *BODY, size=18)
        else:
            slide = _slide(deck, tile["title"])
            _text(slide, tile.get("detail") or "This tile could not be run.", *BODY, size=14, color=INK_700)
    return _bytes(deck)
