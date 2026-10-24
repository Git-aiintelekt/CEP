"""Insights, the report engine's foundation (patch 0049): the definition checked against the catalogue, the SQL the compiler
makes of it, and the API: the catalogue, a report's life, a run under the engine's limits, the export, and what another
workspace or a person without the right cannot do."""
import json

import pytest
from django.conf import settings
from django.db import connections

from datetime import date

from decimal import Decimal
from unittest import mock

from apps.insights import access, ai_report, alerts, analyze, assistant, catalog, comments, slides, clickhouse, compare, compiler, dashboards, definition, derived, drill, embeds, excel, links, pdf, runner, schedules, shape, topn
from apps.insights.definition import DefinitionError

pytestmark = pytest.mark.django_db(databases=["default", "platform"])

P = settings.PLATFORM_ALIAS
DATASET = {
    "key": "messages", "label": "Messages", "source": "reporting.rpt_messages", "workspace_field": "tenant_id", "primary_key": "message_id", "time_field": "submitted_at",
    "fields": [
        {"key": "channel", "label": "Channel", "column": "channel", "type": "string", "role": "dimension", "aggregates": [], "operators": ["=", "!=", "contains", "starts_with", "in", "is_null", "is_not_null"], "buckets": [], "format": None, "sensitivity": "NORMAL"},
        {"key": "campaign_name", "label": "Campaign", "column": "campaign_name", "type": "string", "role": "dimension", "aggregates": [], "operators": ["=", "!=", "contains", "starts_with", "in", "is_null", "is_not_null"], "buckets": [], "format": None, "sensitivity": "NORMAL"},
        {"key": "submitted_at", "label": "Sent", "column": "submitted_at", "type": "datetime", "role": "dimension", "aggregates": [], "operators": ["=", "!=", ">", ">=", "<", "<=", "between", "last_days", "is_null", "is_not_null"], "buckets": ["day", "week", "month", "quarter", "year"], "format": "datetime", "sensitivity": "NORMAL"},
        {"key": "party_id", "label": "Customer", "column": "party_id", "type": "string", "role": "dimension", "aggregates": [], "operators": ["=", "!=", "contains", "starts_with", "in", "is_null", "is_not_null"], "buckets": [], "format": None, "sensitivity": "PII"},
        {"key": "sent", "label": "Sent Count", "column": "sent_n", "type": "number", "role": "measure", "aggregates": ["sum"], "operators": ["=", "!=", ">", ">=", "<", "<=", "between", "in", "is_null", "is_not_null"], "buckets": [], "format": "integer", "sensitivity": "NORMAL"},
        {"key": "messages", "label": "Messages", "column": "message_id", "type": "number", "role": "measure", "aggregates": ["count"], "operators": [">", ">="], "buckets": [], "format": "integer", "sensitivity": "NORMAL"},
    ],
}


def test_definition_and_compiler(check):
    raw = {"columns": [{"field": "campaign_name"}, {"field": "submitted_at", "bucket": "month"}, {"field": "sent", "aggregate": "sum"}],
           "filters": {"combinator": "and", "rules": [{"field": "channel", "operator": "=", "value": "EMAIL"}, {"field": "submitted_at", "operator": "last_days", "value": 30},
                                                      {"combinator": "or", "rules": [{"field": "campaign_name", "operator": "contains", "value": "Fall %"}, {"field": "channel", "operator": "in", "value": ["SMS", "EMAIL"]}]}]},
           "orderBy": [{"field": "sent", "direction": "desc"}], "limit": 50}
    d = definition.validate(DATASET, raw)
    check("a grouped definition: the measure's aggregate marks it, the dimensions become the group", d["grouped"] and [c.get("aggregate") for c in d["columns"]] == [None, None, "sum"])
    c = compiler.compile(DATASET, d, "a0000000-0000-4000-8000-000000000001", "America/New_York", limit=51, offset=0)
    sql = c.sql
    check("the SQL reads the dataset's view, restricts on the workspace, groups by the dimensions and orders by the measure",
          sql.startswith('SELECT "campaign_name" AS "c0", date_trunc(%s, "submitted_at" AT TIME ZONE %s)::date AS "c1", sum("sent_n") AS "c2" FROM "reporting"."rpt_messages" WHERE "tenant_id" = %s')
          and ' GROUP BY 1, 2 ORDER BY' in sql and sql.endswith('ORDER BY "c2" DESC NULLS LAST LIMIT %s OFFSET %s'), sql)
    check("every value is a bound parameter, in the order of the placeholders: the bucket, the zone, the workspace, the filters, then the limit",
          c.params[:3] == ["month", "America/New_York", "a0000000-0000-4000-8000-000000000001"] and c.params[-2:] == [51, 0] and "EMAIL" in c.params and 30 in c.params
          and ["SMS", "EMAIL"] in c.params and sql.count("%s") == len(c.params), str(c.params))
    check("contains escapes the wildcard the person typed", "%Fall \\%%" in c.params and "ILIKE %s ESCAPE" in sql)
    check("a day given for a timestamp is that day in the workspace's zone", "AT TIME ZONE %s) >= now() - make_interval" not in sql and "make_interval(days => %s)" in sql)
    span = compiler.compile(DATASET, definition.validate(DATASET, {"columns": [{"field": "channel"}], "filters": {"rules": [{"field": "submitted_at", "operator": "between", "value": ["2026-09-01", "2026-09-30"]},
                                                                                                                              {"field": "submitted_at", "operator": "=", "value": "2026-10-01"}]}}),
                            "a0000000-0000-4000-8000-000000000001", "America/New_York", limit=10)
    check("a range or a day on a timestamp names the zone once per use of the column, so the placeholders and the parameters agree",
          span.sql.count("%s") == len(span.params) and span.params[1:9] == ["America/New_York", "2026-09-01", "America/New_York", "2026-09-30", "America/New_York", "2026-10-01", "America/New_York", "2026-10-01"], str(span.params))
    check("the result's columns carry the labels, the types and the formats", [x["label"] for x in c.columns] == ["Campaign", "Sent", "Sent Count"] and c.columns[1]["type"] == "date" and c.columns[2]["format"] == "integer")

    def refused(raw, where):
        try:
            definition.validate(DATASET, raw)
            return False
        except DefinitionError as exc:
            return exc.field == where
    check("a measure without an aggregate is refused, naming the column", refused({"columns": [{"field": "sent"}]}, "columns[0].aggregate"))
    check("an unknown field is refused where it sits", refused({"columns": [{"field": "campaign_name"}], "filters": {"rules": [{"field": "nope", "operator": "=", "value": 1}]}}, "filters.rules[0].field"))
    check("an operator a type does not take is refused", refused({"columns": [{"field": "campaign_name"}], "filters": {"rules": [{"field": "sent", "operator": "contains", "value": "x"}]}}, "filters.rules[0].operator"))
    check("a bucket on a string is refused", refused({"columns": [{"field": "campaign_name", "bucket": "month"}]}, "columns[0].bucket"))
    check("sorting by a column the report does not have is refused", refused({"columns": [{"field": "campaign_name"}], "orderBy": [{"field": "sent"}]}, "orderBy[0].field"))
    check("groups nest two levels at most", refused({"columns": [{"field": "campaign_name"}], "filters": {"rules": [{"rules": [{"rules": [{"rules": []}]}]}]}}, "filters.rules[0].rules[0].rules[0]"))
    shown = definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "submitted_at", "bucket": "month"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "pivot", "across": "submitted_at"}})
    check("the view saves with the definition: its type, and the dimension spread across", shown["view"] == {"type": "pivot", "across": "submitted_at"})
    check("spreading a measure, or a field the report does not have, is refused at the view", refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "across": "sent"}}, "view.across")
          and refused({"columns": [{"field": "channel"}], "view": {"type": "tree"}}, "view.type"))
    check("the limit is capped to what the engine allows", definition.validate(DATASET, {"columns": [{"field": "channel"}], "limit": 10**9}, max_limit=200)["limit"] == 200)
    check("the signature is one hash for one normalised definition", definition.signature(d) == definition.signature(definition.validate(DATASET, raw)) and len(definition.signature(d)) == 64)

    # personal data: hidden from a person without the right when building, masked when reading someone else's report
    seen = catalog._visible(DATASET, can_see_pii=False)
    check("without the right the catalogue hides the PII field and remembers it as locked", "party_id" not in [f["key"] for f in seen["fields"]] and seen["locked"] == [{"key": "party_id", "label": "Customer"}])
    try:
        definition.validate(seen, {"columns": [{"field": "party_id"}, {"field": "messages", "aggregate": "count"}]})
        told = ""
    except DefinitionError as exc:
        told = f"{exc.field}: {exc}"
    check("a draft naming it is told why, where it sits", told == "columns[0].field: Customer is personal data: building a report on it needs the report administration right.", told)
    columns = [{"key": "party_id", "label": "Customer", "type": "string", "format": None}, {"key": "messages", "label": "Messages", "type": "number", "format": "integer", "aggregate": "count"}]
    rows = [["p-1", 4], ["p-2", 1], [None, 2]]
    masked_columns, masked_rows = runner.mask(DATASET, columns, rows, True)
    check("a reader without the right gets the PII column's values masked and the column marked; empty values stay empty",
          masked_rows == [[runner.MASK, 4], [runner.MASK, 1], [None, 2]] and masked_columns[0].get("masked") is True and "masked" not in masked_columns[1], str(masked_rows))
    check("a reader with the right gets the rows as they are", runner.mask(DATASET, columns, rows, False) == (columns, rows))

    # the ClickHouse dialect: typed named placeholders, its time-zone and bucket expressions, IN, ILIKE without ESCAPE, FINAL
    ch = compiler.compile(DATASET, d, "a0000000-0000-4000-8000-000000000001", "America/New_York", limit=51, offset=0, dialect="clickhouse")
    check("the ClickHouse query reads the table FINAL, buckets with date_trunc in the zone, binds every value as a typed parameter and keeps the shape",
          ch.sql.startswith('SELECT "campaign_name" AS "c0", toDate(date_trunc(\'month\', toTimeZone("submitted_at", {p0:String}))) AS "c1", sum("sent_n") AS "c2" FROM "reporting"."rpt_messages" FINAL WHERE "tenant_id" = {p1:UUID}')
          and "toIntervalDay({p" in ch.sql and " IN {p" in ch.sql and "ESCAPE" not in ch.sql and "ILIKE {p" in ch.sql and "LIMIT {p" in ch.sql and isinstance(ch.params, dict)
          and ch.params["p0"] == "America/New_York" and ch.params["p1"] == "a0000000-0000-4000-8000-000000000001" and ["SMS", "EMAIL"] in ch.params.values(), ch.sql)
    ch_span = compiler.compile(DATASET, definition.validate(DATASET, {"columns": [{"field": "channel"}], "filters": {"rules": [{"field": "submitted_at", "operator": "between", "value": ["2026-09-01", "2026-09-30"]}]}}),
                               "a0000000-0000-4000-8000-000000000001", "America/New_York", limit=10, dialect="clickhouse")
    check("a range of days on a timestamp becomes the instants of the local days' bounds in ClickHouse", '"submitted_at" >= toDateTime({p1:String}, {p2:String}) AND "submitted_at" < toDateTime({p3:String}, {p4:String}) + INTERVAL 1 DAY' in ch_span.sql
          and ch_span.params["p1"] == "2026-09-01" and ch_span.params["p2"] == "America/New_York", ch_span.sql)
    ch_t = compiler.compile_totals(DATASET, d2 if "d2" in dir() else d, "a0000000-0000-4000-8000-000000000001", "America/New_York", [(), ("campaign_name",)], limit=100, dialect="clickhouse")
    check("the totals query keeps GROUPING SETS and GROUPING in ClickHouse", "GROUP BY GROUPING SETS ((), (\"d0\"))" in ch_t.sql and 'GROUPING("d0") AS "g"' in ch_t.sql and " FINAL WHERE " in ch_t.sql, ch_t.sql)

    # the ClickHouse connection: a back-office profile row, else the server environment, else none
    from django.test import override_settings
    with override_settings(CLICKHOUSE_URL="", FIELD_ENCRYPTION_KEY=settings.FIELD_ENCRYPTION_KEY):
        check("no profile and no environment: no connection", clickhouse.from_environment() is None and clickhouse.test(None)[0] is False)
    with override_settings(CLICKHOUSE_URL="http://ch.example:8123", CLICKHOUSE_USER="reader", CLICKHOUSE_PASSWORD="x", CLICKHOUSE_DATABASE="cep_reporting"):
        env = clickhouse.from_environment()
        check("the environment makes a connection with its database and no profile", env.url == "http://ch.example:8123" and env.user == "reader" and env.database == "cep_reporting" and env.source == "the server environment")
    from apps.platform_ops import crypto
    row = {"profile": "analytics-eu", "url": "http://eu.example:8123", "username": "cep", "password_enc": crypto.encrypt("s3cret"), "database_name": "cep_eu"}
    named = clickhouse.from_row(row)
    check("a profile row makes a connection with its password decrypted and its own database", named.password == "s3cret" and named.database == "cep_eu" and named.source == "profile analytics-eu"
          and clickhouse.table("messages", named) == "cep_eu.rpt_messages")

    # the engine's totals: the grouping sets a view needs, and their query
    d2 = definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "submitted_at", "bucket": "month"}, {"field": "sent", "aggregate": "sum"}, {"field": "messages", "aggregate": "count"}],
                                       "filters": {"rules": [{"field": "channel", "operator": "=", "value": "EMAIL"}]}})
    check("a hierarchy totals every prefix of its dimensions and the whole; a pivot its rows, its columns and the whole; a grid nothing",
          compiler.totals_sets(d2, {"type": "hierarchy"}) == [(), ("channel",)] and compiler.totals_sets(d2, {"type": "pivot", "across": "submitted_at"}) == [(), ("channel",), ("submitted_at",)]
          and compiler.totals_sets(d2, {"type": "grid"}) is None and compiler.totals_sets(d2, {"type": "pivot"}) is None)
    t = compiler.compile_totals(DATASET, d2, "a0000000-0000-4000-8000-000000000001", "America/New_York", [(), ("channel",), ("submitted_at",)], limit=5000)
    check("the totals query computes the dimensions once in a subquery and totals them by GROUPING SETS, with GROUPING last, every placeholder bound",
          t.sql.startswith('SELECT "d0", "d1", sum("m0") AS "v0", count("m1") AS "v1", GROUPING("d0", "d1") AS "g" FROM (SELECT "channel" AS "d0", date_trunc(%s, "submitted_at" AT TIME ZONE %s)::date AS "d1", "sent_n" AS "m0", "message_id" AS "m1" FROM')
          and 'GROUP BY GROUPING SETS ((), ("d0"), ("d1")) LIMIT %s' in t.sql and t.sql.count("%s") == len(t.params) and t.params[-1] == 5000, t.sql)
    check("exact totals win over combining the rows: an average subtotal comes from the engine",
          shape.hierarchy([{"key": "channel", "label": "Channel", "type": "string", "format": None}, {"key": "submitted_at", "label": "Sent", "type": "date", "format": "datetime", "bucket": "month"},
                           {"key": "opened", "label": "Opened", "type": "number", "format": "integer", "aggregate": "avg"}],
                          [["EMAIL", "2026-09-01", 10.5], ["EMAIL", "2026-08-01", 9.0]], totals=[{"keys": {"channel": "EMAIL"}, "values": [9.9]}, {"keys": {}, "values": [9.9]}])["rows"][0][2] == ["EMAIL", 9.9])

    # the pivot and the hierarchy as tables (the Excel export), and the workbook itself
    cols = [{"key": "channel", "label": "Channel", "type": "string", "format": None}, {"key": "submitted_at", "label": "Sent", "type": "date", "format": "datetime", "bucket": "month"},
            {"key": "sent", "label": "Sent Count", "type": "number", "format": "integer", "aggregate": "sum"}, {"key": "opened", "label": "Opened", "type": "number", "format": "integer", "aggregate": "avg"}]
    data = [["EMAIL", "2026-09-01", 356920, 10.5], ["EMAIL", "2026-08-01", 247250, 9.0], ["SMS", "2026-09-01", 29580, 0], ["SMS", "2026-08-01", 27810, 0]]
    pv = shape.pivot(cols, data, "submitted_at")
    check("a pivot spreads the dimension across with a total on both edges, in the rows' order", pv["header"] == ["Channel", "2026-09-01", "2026-08-01", "Total"] and pv["rows"] == [["EMAIL", 356920, 247250, 604170], ["SMS", 29580, 27810, 57390]]
          and pv["total"] == ["Total", 386500, 275060, 661560], str(pv.get("rows")))
    hy = shape.hierarchy(cols, data)
    check("a hierarchy folds the dimensions in order with a subtotal per group; an average cannot be combined",
          [(d, g, cells[0]) for d, g, cells in hy["rows"]] == [(0, True, "EMAIL"), (1, False, "2026-09-01"), (1, False, "2026-08-01"), (0, True, "SMS"), (1, False, "2026-09-01"), (1, False, "2026-08-01")]
          and hy["rows"][0][2][1:] == [604170, None] and hy["total"] == ["Total", 661560, None], str(hy["rows"]))
    check("a pivot without a dimension across, or a hierarchy on one dimension, says what it needs", "reason" in shape.pivot(cols, data, None) and "reason" in shape.hierarchy(cols[1:], [r[1:] for r in data]))
    import zipfile
    book = excel.build("Monthly Sends", cols, data, {"type": "pivot", "across": "submitted_at"})
    with zipfile.ZipFile(__import__("io").BytesIO(book)) as z:
        names = z.namelist()
        strings = z.read("xl/sharedStrings.xml").decode()
        sheets = z.read("xl/workbook.xml").decode()
    check("the workbook has the pivot sheet first and the rows after, with the labels and the values as text", "xl/worksheets/sheet1.xml" in names and "xl/worksheets/sheet2.xml" in names
          and "Pivot" in sheets and "Sent Count" in strings and "EMAIL" in strings, str(names))

    # the PDF: a report in each view, and a dashboard with its tiles, as real PDF files
    for kind, across in (("grid", None), ("pivot", "submitted_at"), ("hierarchy", None), ("chart", "submitted_at")):
        doc = pdf.report("Monthly Sends", cols, data, {"type": kind, "across": across}, workspace="Northwind Apparel", when="16 Oct 2026")
        check(f"the {kind} PDF is a PDF", doc[:5] == b"%PDF-" and len(doc) > 1500, str(len(doc)))
    board = pdf.dashboard("Marketing Weekly", "Sends and revenue", "16 Sep – 15 Oct 2026", [
        {"title": "Sent", "kind": "figure", "value": 391520, "previous": 361180, "compared": True, "column": cols[2]},
        {"title": "Sends by Month", "kind": "view", "view": {"type": "chart", "across": "channel"}, "result": {"columns": cols, "rows": data, "row_count": 4}, "own_dates": False},
        {"title": "Gone", "kind": "missing", "detail": "The report behind this tile is gone."}], workspace="Northwind Apparel", when="16 Oct 2026")
    check("the dashboard PDF is a PDF with its figures, views and a missing tile", board[:5] == b"%PDF-" and len(board) > 2000, str(len(board)))

    # scheduled delivery: when the next sending falls, in the workspace's zone
    from datetime import datetime, time as clock
    from zoneinfo import ZoneInfo
    ny = ZoneInfo("America/New_York")
    after = datetime(2026, 10, 16, 12, 0, tzinfo=ny)                                 # a Friday, noon
    nxt = lambda every, on_day, at: schedules.next_run(every, on_day, clock(*at), ny, after).astimezone(ny).strftime("%a %Y-%m-%d %H:%M")
    check("every day at 08:00 after Friday noon is Saturday 08:00; a weekday is Monday; weekly on Monday is Monday; monthly on the 1st is 1 November",
          nxt("day", None, (8, 0)) == "Sat 2026-10-17 08:00" and nxt("weekday", None, (7, 0)) == "Mon 2026-10-19 07:00" and nxt("week", 1, (8, 0)) == "Mon 2026-10-19 08:00"
          and nxt("month", 1, (9, 0)) == "Sun 2026-11-01 09:00" and nxt("day", None, (18, 0)) == "Fri 2026-10-16 18:00", nxt("month", 1, (9, 0)))
    check("a schedule reads in words", schedules.words({"every": "week", "on_day": 1, "at_time": clock(8, 0)}) == "Every Monday at 08:00"
          and schedules.words({"every": "month", "on_day": 22, "at_time": clock(9, 30)}) == "Monthly on the 22nd at 09:30" and schedules.words({"every": "weekday", "on_day": None, "at_time": "07:00:00"}) == "Every weekday at 07:00")

    # dashboards: the period's bounds, and the period as one more rule on the time field
    check("this month runs from the 1st to today, the previous period is the same length just before",
          dashboards.period_bounds({"preset": "month"}, date(2026, 10, 15)) == {"from": "2026-10-01", "to": "2026-10-15", "previous_from": "2026-09-16", "previous_to": "2026-09-30"})
    check("the quarter starts on its first day", dashboards.period_bounds({"preset": "quarter"}, date(2026, 11, 3))["from"] == "2026-10-01")
    shaped = dashboards._with_period(raw, "submitted_at", "2026-09-01", "2026-09-30")
    check("the period is and-ed onto the report's own rules, and a report with its own rule on the time field is left alone",
          shaped["filters"]["rules"][-1] == {"field": "submitted_at", "operator": "between", "value": ["2026-09-01", "2026-09-30"]} and len(shaped["filters"]["rules"]) == 4
          and dashboards._has_time_rule(raw["filters"], "submitted_at") and not dashboards._has_time_rule({"rules": [{"field": "channel", "operator": "=", "value": "x"}]}, "submitted_at"))

    # calculated measures: a formula over other measures, compiled in the select and in the totals, checked against the dataset
    T = "a0000000-0000-4000-8000-000000000001"
    RATE = {"key": "delivered_rate", "label": "Delivered Rate", "column": None, "type": "number", "role": "measure", "aggregates": ["calc"], "operators": [], "buckets": [], "format": "percent", "sensitivity": "NORMAL",
            "formula": {"op": "ratio", "left": {"field": "sent", "aggregate": "sum"}, "right": {"field": "messages", "aggregate": "count"}, "scale": 100}}
    calc_ds = {**DATASET, "fields": [*DATASET["fields"], RATE]}
    dc = definition.validate(calc_ds, {"columns": [{"field": "channel"}, {"field": "delivered_rate", "aggregate": "calc"}]})
    cc = compiler.compile(calc_ds, dc, T, "America/New_York", limit=10)
    check("a calculated measure compiles to its formula over the operands' aggregates, dividing by NULLIF so an empty denominator reads as empty",
          '((sum("sent_n")) * 1.0 / NULLIF((count("message_id")), 0)) * 100.0 AS "c1"' in cc.sql and cc.columns[1] == {"key": "delivered_rate", "label": "Delivered Rate", "type": "number", "format": "percent", "aggregate": "calc"}, cc.sql)
    tc = compiler.compile_totals(calc_ds, dc, T, "America/New_York", [(), ("channel",)], limit=100)
    check("the totals query carries each operand's raw column once and totals the formula in the outer query",
          '"sent_n" AS "m0_0", "message_id" AS "m0_1"' in tc.sql and '((sum("m0_0")) * 1.0 / NULLIF((count("m0_1")), 0)) * 100.0 AS "v0"' in tc.sql, tc.sql)
    check("MySQL backticks the operands; ClickHouse writes nullIf",
          '(sum(`sent_n`)) * 1.0 / NULLIF((count(`message_id`)), 0)' in compiler.compile(calc_ds, dc, T, "America/New_York", limit=10, dialect="mysql").sql
          and 'nullIf((count("message_id")), 0)' in compiler.compile(calc_ds, dc, T, "America/New_York", limit=10, dialect="clickhouse").sql)
    diff = {**RATE, "key": "net", "label": "Net", "formula": {"op": "difference", "left": {"field": "sent", "aggregate": "sum"}, "right": {"value": 10}}}
    nested = {**RATE, "key": "net_rate", "label": "Net Rate", "formula": {"op": "ratio", "left": {"field": "net", "aggregate": "calc"}, "right": {"field": "sent", "aggregate": "sum"}}}
    nested_ds = {**DATASET, "fields": [*DATASET["fields"], diff, nested]}
    nc = compiler.compile(nested_ds, definition.validate(nested_ds, {"columns": [{"field": "net_rate", "aggregate": "calc"}]}), T, "America/New_York", limit=1)
    check("a formula may read another calculated measure and a number", '((sum("sent_n")) - (10.0)) * 1.0 / NULLIF((sum("sent_n")), 0) AS "c0"' in nc.sql, nc.sql)
    check("a formula is checked against the dataset's fields: an operand that is not a measure, an aggregate it does not take, a side that is neither a field nor a number, an unknown op",
          catalog.formula_error(calc_ds["fields"], RATE["formula"]) is None
          and catalog.formula_error(calc_ds["fields"], {"op": "ratio", "left": {"field": "channel", "aggregate": "sum"}, "right": {"value": 2}}) == "The left side names a measure of this dataset."
          and catalog.formula_error(calc_ds["fields"], {"op": "ratio", "left": {"field": "sent", "aggregate": "avg"}, "right": {"value": 2}}) == "Sent Count does not take avg."
          and catalog.formula_error(calc_ds["fields"], {"op": "ratio", "left": {"field": "sent", "aggregate": "sum"}, "right": {"value": "x"}}) == "The right side names a measure or a number."
          and catalog.formula_error(calc_ds["fields"], {"op": "modulo", "left": {}, "right": {}}) is not None)
    customers = {"key": "customers", "label": "Customers", "column": "party_id", "type": "number", "role": "measure", "aggregates": ["count_distinct"], "operators": [], "buckets": [], "format": "integer", "sensitivity": "PII"}
    per_customer = {**RATE, "key": "per_customer", "label": "Sends per Customer", "format": "decimal", "formula": {"op": "ratio", "left": {"field": "sent", "aggregate": "sum"}, "right": {"field": "customers", "aggregate": "count_distinct"}}}
    seen_keys = [f["key"] for f in catalog._visible({**DATASET, "fields": [*DATASET["fields"], customers, per_customer, RATE]}, can_see_pii=False)["fields"]]
    check("a calculated measure over personal data goes with it for a reader without the right; the others stay", "per_customer" not in seen_keys and "customers" not in seen_keys and "delivered_rate" in seen_keys, str(seen_keys))

    # pie and funnel are chart kinds; the PDF draws them
    shown = definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "chart": "pie"}})
    check("pie and funnel are chart kinds, and nothing else is", shown["view"] == {"type": "chart", "chart": "pie"} and refused({"columns": [{"field": "channel"}], "view": {"type": "chart", "chart": "donut"}}, "view.chart"))
    for kind in ("pie", "funnel"):
        doc = pdf.report("Sends by Channel", cols, data, {"type": "chart", "chart": kind}, workspace="Northwind Apparel", when="16 Oct 2026")
        check(f"the {kind} PDF is a PDF", doc[:5] == b"%PDF-" and len(doc) > 1500, str(len(doc)))

    # dashboard slicers: the request's values become IN rules on the datasets that have the field, and nothing on the others
    rules = dashboards._slicer_rules(DATASET, {"channel": ["EMAIL", "SMS"], "store": ["Downtown"], "sent": [1]})
    check("a slicer becomes an IN rule on a dimension the dataset has; a field it lacks, or a measure, is skipped", rules == [{"field": "channel", "operator": "in", "value": ["EMAIL", "SMS"]}], str(rules))
    check("the request's slicers are read as fields to lists of scalars", dashboards._slicers_body({"channel": ["EMAIL", {"x": 1}], "empty": [], "n": "x"}) == {"channel": ["EMAIL"]})
    check("a slicer in the layout is a text, yes/no or number dimension of the tiles' datasets, labelled from the field",
          dashboards._slicers_layout([{"field": "channel"}], [DATASET]) == [{"field": "channel", "label": "Channel"}])
    for bad in ([{"field": "sent"}], [{"field": "submitted_at"}], [{"field": "channel"}, {"field": "channel"}]):
        try:
            dashboards._slicers_layout(bad, [DATASET])
            check("a measure, a date or a repeated field is refused as a slicer", False, str(bad))
        except Exception as exc:                                                       # noqa: BLE001
            check("a measure, a date or a repeated field is refused as a slicer", "layout.slicers" in str(exc), str(exc))
    # time comparisons: checked with the definition, worked out from the rows and the shifted run
    timed = definition.validate(DATASET, {"columns": [{"field": "submitted_at", "bucket": "month"}, {"field": "channel"},
                                                      {"field": "sent", "aggregate": "sum", "compare": ["previous_period", "change_previous_year", "running_total", "percent_of_total"]}],
                                          "filters": {"combinator": "and", "rules": [{"field": "submitted_at", "operator": "between", "value": ["2026-07-01", "2026-09-30"]}]}})
    check("a measure keeps its comparisons, in order", timed["columns"][2]["compare"] == ["previous_period", "change_previous_year", "running_total", "percent_of_total"])
    check("comparing with an earlier period needs a date in the report", refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum", "compare": ["previous_year"]}]}, "columns[1].compare"))
    check("a share of the total needs a sum or a count, and an unknown or repeated comparison is refused",
          refused({"columns": [{"field": "channel"}, {"field": "messages", "aggregate": "count", "compare": ["percent_of_total", "percent_of_total"]}]}, "columns[1].compare")
          and refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum", "compare": ["yesterday"]}]}, "columns[1].compare")
          and refused({"columns": [{"field": "channel", "compare": ["percent_of_total"]}]}, "columns[0].compare"))
    check("the time axis is the first bucketed date", compare.axis(DATASET, timed["columns"]) == (0, "submitted_at", "month"))
    moved = compare.shifted(timed, "submitted_at", 12, 0, date(2026, 10, 17))
    check("a year back, the date rule moves and the last day of a month stays the last day",
          moved["filters"]["rules"] == [{"field": "submitted_at", "operator": "between", "value": ["2025-07-01", "2025-09-30"]}]
          and compare.shifted(timed, "submitted_at", 1, 0, date(2026, 10, 17))["filters"]["rules"][0]["value"] == ["2026-06-01", "2026-08-31"], str(moved["filters"]))
    check("in the last 30 days becomes the same days a year earlier",
          compare.shifted({"filters": {"combinator": "and", "rules": [{"field": "submitted_at", "operator": "last_days", "value": 30}]}}, "submitted_at", 12, 0, date(2026, 10, 17))["filters"]["rules"][0]["value"] == ["2025-09-17", "2025-10-17"])
    check("no rule on the date: the rows serve as they are; a date rule inside an or-group cannot be moved",
          compare.shifted({"filters": {"combinator": "and", "rules": [{"field": "channel", "operator": "=", "value": "SMS"}]}}, "submitted_at", 12, 0, date(2026, 10, 17)) is None)
    try:
        compare.shifted({"filters": {"combinator": "or", "rules": [{"field": "submitted_at", "operator": "last_days", "value": 7}, {"field": "channel", "operator": "=", "value": "SMS"}]}}, "submitted_at", 12, 0, date(2026, 10, 17))
        check("a date rule inside an or-group is refused", False)
    except compare.CompareError:
        check("a date rule inside an or-group is refused", True)
    rows = [["2026-07-01", "EMAIL", 100], ["2026-08-01", "EMAIL", 150], ["2026-09-01", "EMAIL", 120], ["2026-08-01", "SMS", 50], ["2026-09-01", "SMS", 80]]
    june = [["2026-06-01", "EMAIL", 80], *rows]                                           # the shifted run (a month back) reaches June
    last_year = [["2025-07-01", "EMAIL", 50], ["2025-09-01", "SMS", 0]]
    out = compare.apply(DATASET, timed, rows, {"period": june, "year": last_year}, {2: 500.0})
    check("the previous month comes from the shifted run, so July has June's figure", [r[3] for r in out] == [80, 100, 150, None, 50], str(out))
    check("the change against last year is in percent, empty where last year had nothing to divide by", out[0][4] == 100.0 and out[4][4] is None and out[1][4] is None, str([r[4] for r in out]))
    check("the running total adds up along the dates per channel", [r[5] for r in out] == [100.0, 250.0, 370.0, 50.0, 130.0], str([r[5] for r in out]))
    check("the share of the total is of the report's whole", [round(r[6], 1) for r in out] == [20.0, 30.0, 24.0, 10.0, 16.0])
    cols = compare.columns(DATASET, timed, [{"key": "submitted_at", "label": "Sent", "type": "date", "format": "datetime", "bucket": "month"}, {"key": "channel", "label": "Channel", "type": "string", "format": None},
                                            {"key": "sent", "label": "Sent Count", "type": "number", "format": "integer", "aggregate": "sum"}])
    check("each comparison is a column after its measure, labelled in words, a change and a share as percentages",
          [c["label"] for c in cols[2:]] == ["Sent Count", "Sent Count, Previous Month", "Sent Count, Change vs Last Year", "Sent Count, Running Total", "Sent Count, % of Total"]
          and [c.get("compare") for c in cols[3:]] == ["previous_period", "change_previous_year", "running_total", "percent_of_total"] and cols[4]["format"] == "percent" and cols[3]["format"] == "integer")
    check("the engine's totals line up with the measure columns, empty for each comparison",
          compare.expand_totals(timed, [{"keys": {}, "values": [500]}]) == [{"keys": {}, "values": [500, None, None, None, None]}])
    check("a subtotal never adds up a running total or a share", shape._agg({"aggregate": "sum", "compare": "running_total"}) is None and shape._agg({"aggregate": "sum"}) == "sum")
    check("a week compares with the same week 52 weeks earlier; a month with the month before", compare.offset("year", "week") == (0, 364) and compare.offset("period", "month") == (1, 0)
          and compare.add_months(date(2024, 2, 29), 12) == date(2025, 2, 28))

    # conditional formatting and the new chart kinds
    formatted = definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum", "compare": ["percent_of_total"]}, {"field": "messages", "aggregate": "count"}],
                                              "view": {"type": "grid", "formats": [{"column": "sent", "style": "bars"}, {"column": "sent__percent_of_total", "style": "rules",
                                                                                                                     "rules": [{"op": "<", "value": 10, "tone": "bad"}, {"op": ">=", "value": 50, "tone": "good"}]},
                                                                                    {"column": "messages", "style": "scale"}]}})
    check("formatting saves with the view: bars, rules on a comparison column, a colour scale",
          [f["style"] for f in formatted["view"]["formats"]] == ["bars", "rules", "scale"] and formatted["view"]["formats"][1]["rules"][1] == {"op": ">=", "value": 50, "tone": "good"})
    check("a dimension, a column twice, a rule without a number or with an unknown tone are refused where they sit",
          refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "grid", "formats": [{"column": "channel", "style": "scale"}]}}, "view.formats[0].column")
          and refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "grid", "formats": [{"column": "sent", "style": "bars"}, {"column": "sent", "style": "scale"}]}}, "view.formats[1].column")
          and refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "grid", "formats": [{"column": "sent", "style": "rules", "rules": [{"op": "<", "value": "ten", "tone": "bad"}]}]}}, "view.formats[0].rules[0]")
          and refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "grid", "formats": [{"column": "sent", "style": "rules", "rules": [{"op": "<", "value": 1, "tone": "red"}]}]}}, "view.formats[0].rules[0]"))
    sample_cols = [{"key": "channel", "label": "Channel", "type": "string", "format": None}, {"key": "submitted_at", "label": "Sent", "type": "date", "format": "datetime", "bucket": "month"},
                   {"key": "sent", "label": "Sent Count", "type": "number", "format": "integer", "aggregate": "sum"}, {"key": "opened", "label": "Opened", "type": "number", "format": "integer", "aggregate": "sum"}]
    sample_rows = [["EMAIL", "2026-09-01", 356920, 134210], ["EMAIL", "2026-08-01", 247250, 91880], ["SMS", "2026-09-01", 29580, 0], ["SMS", "2026-08-01", 27810, 0]]
    for kind in ("stacked", "stacked100", "combo"):
        check(f"{kind} is a chart kind", definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "chart": kind}})["view"]["chart"] == kind)
        doc = pdf.report("Sends by Month", sample_cols, sample_rows, {"type": "chart", "chart": kind, "across": "channel"}, workspace="Northwind Apparel", when="18 Oct 2026")
        check(f"the {kind} PDF is a PDF", doc[:5] == b"%PDF-" and len(doc) > 1500, str(len(doc)))

    # the time measures: year and quarter to date, a moving average, a rank, a target
    measured = definition.validate(DATASET, {"columns": [{"field": "submitted_at", "bucket": "month"}, {"field": "channel"},
                                                         {"field": "sent", "aggregate": "sum", "compare": ["year_to_date", "quarter_to_date", "moving_average", "rank", "percent_of_target", "difference_to_target"], "target": 100}],
                                             "filters": {"combinator": "and", "rules": [{"field": "submitted_at", "operator": "between", "value": ["2026-08-01", "2026-09-30"]}]}})
    check("a measure takes up to six comparisons and keeps its target", len(measured["columns"][2]["compare"]) == 6 and measured["columns"][2]["target"] == 100)
    check("a target comparison needs the target; year to date adds up sums and counts only",
          refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum", "compare": ["percent_of_target"]}]}, "columns[1].target")
          and refused({"columns": [{"field": "submitted_at", "bucket": "month"}, {"field": "messages", "aggregate": "count", "compare": ["rank"]}, {"field": "sent", "aggregate": "sum", "compare": ["year_to_date"]}]}, "") is False)
    wide = compare.widened(measured, "submitted_at", compare.reach_back(["year_to_date", "moving_average"], "month"), date(2026, 10, 17))
    check("to reach back, the lower bound moves to the year's start or two months back, whichever is earlier; the upper bound stays",
          wide["filters"]["rules"][0]["value"] == ["2026-01-01", "2026-09-30"] and compare.reach_back(["moving_average"], "month")(date(2026, 8, 1)) == date(2026, 6, 1))
    rows_m = [["2026-08-01", "EMAIL", 150], ["2026-09-01", "EMAIL", 120], ["2026-08-01", "SMS", 50], ["2026-09-01", "SMS", 80]]
    reached = [["2026-01-01", "EMAIL", 10], ["2026-06-01", "EMAIL", 30], ["2026-07-01", "EMAIL", 100], *rows_m, ["2026-07-01", "SMS", 20]]
    out_m = compare.apply(DATASET, measured, rows_m, {}, None, reached)
    ytd, qtd, mavg, rank, pct, diff = ([r[i] for r in out_m] for i in range(3, 9))
    check("year to date adds every month of the year so far, from before the report's first month", ytd == [290.0, 410.0, 70.0, 150.0], str(ytd))
    check("quarter to date restarts at July and October", qtd == [250.0, 370.0, 70.0, 150.0], str(qtd))
    check("the moving average is of the month and the two before, a missing month counting as nothing", [round(v, 1) for v in mavg] == [93.3, 123.3, 23.3, 50.0], str(mavg))
    check("the rank is per month among the channels, the largest first", rank == [1, 1, 2, 2], str(rank))
    check("against a target of 100: the percentage and the difference", pct == [150.0, 120.0, 50.0, 80.0] and diff == [50.0, 20.0, -50.0, -20.0], str((pct, diff)))
    labels = [c["label"] for c in compare.columns(DATASET, measured, [{"key": "submitted_at", "label": "Sent", "type": "date", "format": "datetime", "bucket": "month"},
                                                                        {"key": "channel", "label": "Channel", "type": "string", "format": None},
                                                                        {"key": "sent", "label": "Sent Count", "type": "number", "format": "integer", "aggregate": "sum"}])[3:]]
    check("the new columns are named in words", labels == ["Sent Count, Year to Date", "Sent Count, Quarter to Date", "Sent Count, Moving Average (3 Months)", "Sent Count, Rank",
                                                           "Sent Count, % of Target", "Sent Count, Difference to Target"], str(labels))
    flat_rank = compare.apply(DATASET, definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum", "compare": ["rank"]}]}), [["EMAIL", 5], ["SMS", 9], ["PUSH", 5]], {}, None)
    check("without a date, the rank is across the whole report, ties sharing a place", [r[2] for r in flat_rank] == [2, 1, 2])

    # drill-through: the fields a report takes from others, checked against the dataset
    detail = definition.validate(DATASET, {"columns": [{"field": "campaign_name"}, {"field": "sent", "aggregate": "sum"}], "drillthrough": ["channel", "channel", "submitted_at"]})
    check("a detail report names the dimensions it takes from other reports, each once", detail["drillthrough"] == ["channel", "submitted_at"])
    check("a measure, or an unknown field, cannot be taken", refused({"columns": [{"field": "channel"}], "drillthrough": ["sent"]}, "drillthrough[0]")
          and refused({"columns": [{"field": "channel"}], "drillthrough": ["nope"]}, "drillthrough[0]"))

    # row access: a person's rules and-ed onto the report's own filters, typed as the field
    original = access.rules_for
    try:
        access.rules_for = lambda tenant_id, user_id, key: {"channel": ["SMS", "EMAIL"]}
        base = definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "filters": {"combinator": "or", "rules": [{"field": "channel", "operator": "=", "value": "PUSH"}]}})
        limited, notes = access.restrict(T, type("U", (), {"pk": "u1"})(), DATASET, base)
        check("a rule adds an IN rule beside the report's own filters, never inside its or-group", limited["filters"] == {"combinator": "and", "rules": [base["filters"], {"field": "channel", "operator": "in", "value": ["SMS", "EMAIL"]}]}
              and notes == [{"field": "channel", "label": "Channel", "values": ["SMS", "EMAIL"]}], str(limited["filters"]))
        lc = compiler.compile(DATASET, limited, T, "America/New_York", limit=10)
        check("the compiled query keeps the rule", "= ANY(%s)" in lc.sql and ["SMS", "EMAIL"] in lc.params, lc.sql)
        access.rules_for = lambda tenant_id, user_id, key: {"gone_field": ["x"]}
        try:
            access.restrict(T, type("U", (), {"pk": "u1"})(), DATASET, base)
            check("a rule on a field the dataset no longer offers refuses the run", False)
        except access.AccessError:
            check("a rule on a field the dataset no longer offers refuses the run", True)
        access.rules_for = lambda tenant_id, user_id, key: {}
        check("no rule, nothing added", access.restrict(T, type("U", (), {"pk": "u1"})(), DATASET, base) == (base, []))
    finally:
        access.rules_for = original
    where_sql, where_args = access.visible_sql("report", "r", "u1", False)
    check("a report is visible to the workspace, its owner, or a share reaching the person or one of their roles",
          "r.visibility = 'WORKSPACE'" in where_sql and "r.owner_user_id = %s" in where_sql and (not access.present("report_share") or "user_role" in where_sql) and where_args[0] == "u1")
    check("an administrator sees everything", access.visible_sql("report", "r", "u1", True) == ("TRUE", []))

    # custom fields: groups, bands and if-then labels, as CASE over the sources, usable as rows and filters
    custom = definition.validate(DATASET, {"columns": [{"field": "c_region"}, {"field": "c_size"}, {"field": "sent", "aggregate": "sum"}],
                                           "derived": [{"label": "Region", "kind": "group", "field": "campaign_name", "groups": [{"label": "North", "values": ["Fall Launch", "O'Brien's Sale"]}]},
                                                       {"label": "Size", "kind": "band", "field": "sent", "edges": [100, 1000]},
                                                       {"label": "Tier", "kind": "rule", "rules": [{"field": "sent", "op": ">=", "value": 500, "label": "Big"}], "other": "Small"}],
                                           "filters": {"combinator": "and", "rules": [{"field": "c_tier", "operator": "=", "value": "Big"}]}})
    check("custom fields get keys from their labels, bands their words", [d["key"] for d in custom["derived"]] == ["c_region", "c_size", "c_tier"]
          and custom["derived"][1]["labels"] == ["Under 100", "100 to under 1,000", "1,000 and over"])
    cc2 = compiler.compile(DATASET, custom, T, "America/New_York", limit=10)
    check("a group is a CASE over its source with the values escaped as literals; a band compares with its edges; a filter on a custom field reads its CASE",
          """(CASE WHEN "campaign_name" IN ('Fall Launch', 'O''Brien''s Sale') THEN 'North' ELSE 'Other' END) AS """ + '"c0"' in cc2.sql
          and """(CASE WHEN "sent_n" IS NULL THEN NULL WHEN "sent_n" < 100 THEN 'Under 100' WHEN "sent_n" < 1000 THEN '100 to under 1,000' ELSE '1,000 and over' END) AS """ + '"c1"' in cc2.sql
          and """(CASE WHEN "sent_n" >= 500 THEN 'Big' ELSE 'Small' END) = %s""" in cc2.sql and cc2.sql.count("%s") == len(cc2.params), cc2.sql)
    check("a percent sign in a value is doubled for the positional dialects, a backslash too for MySQL",
          derived._literal("50% off", "postgresql") == "'50%% off'" and derived._literal("a\\b", "mysql") == "'a\\\\b'" and derived._literal("50% off", "clickhouse") == "'50% off'")
    check("a custom field over personal data is personal, and a band needs a number source, edges going up",
          derived.augment(DATASET, {"derived": [{"key": "c_who", "label": "Who", "kind": "group", "field": "party_id", "groups": [{"label": "A", "values": ["p1"]}], "other": "Other"}]})["fields"][-1]["sensitivity"] == "PII"
          and refused({"columns": [{"field": "channel"}], "derived": [{"label": "S", "kind": "band", "field": "channel", "edges": [1]}]}, "derived[0].field")
          and refused({"columns": [{"field": "channel"}], "derived": [{"label": "S", "kind": "band", "field": "sent", "edges": [10, 5]}]}, "derived[0].edges"))
    tc2 = compiler.compile_totals(DATASET, custom, T, "America/New_York", [(), ("c_region",)], limit=10)
    check("the totals query computes a custom field in its subquery too", """(CASE WHEN "campaign_name" IN""" in tc2.sql and tc2.sql.count("%s") == len(tc2.params))

    # Top N with Others: the first dimension's top values by the sorted measure, the rest folded per other dimension
    ranked = definition.validate(DATASET, {"columns": [{"field": "campaign_name"}, {"field": "channel"}, {"field": "sent", "aggregate": "sum"}, {"field": "messages", "aggregate": "count"}],
                                           "orderBy": [{"field": "sent", "direction": "desc"}], "top": {"n": 2, "others": True}})
    tcols = [{"key": "campaign_name"}, {"key": "channel"}, {"key": "sent", "aggregate": "sum"}, {"key": "messages", "aggregate": "count"}]
    trows = [["A", "EMAIL", 50, 5], ["B", "EMAIL", 40, 4], ["C", "EMAIL", 30, 3], ["D", "SMS", 20, 2], ["C", "SMS", 1, 1], ["A", "SMS", 5, 1]]
    out_t = topn.apply(ranked, tcols, trows)
    check("Top 2 keeps A and B; C and D fold into Others per channel, summed", out_t == [["A", "EMAIL", 50, 5], ["B", "EMAIL", 40, 4], ["A", "SMS", 5, 1], ["Others", "EMAIL", 30, 3], ["Others", "SMS", 21, 3]], str(out_t))
    check("without Others the rest is left out; top needs a value to rank by",
          topn.apply({**ranked, "top": {"n": 1, "others": False}}, tcols, trows) == [["A", "EMAIL", 50, 5], ["A", "SMS", 5, 1]]
          and refused({"columns": [{"field": "channel"}], "top": {"n": 3}}, "top") and refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "top": {"n": 0}}, "top"))
    check("an average or a comparison in an Others line is empty", topn.apply({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "avg"}], "top": {"n": 1, "others": True}},
                                                                             [{"key": "channel"}, {"key": "sent", "aggregate": "avg"}], [["A", 5], ["B", 3]]) == [["A", 5], ["Others", None]])

    # drilling in a pivot or a hierarchy keeps the levels under the one drilled
    stacked_levels = definition.validate(DATASET, {"columns": [{"field": "submitted_at", "bucket": "year"}, {"field": "channel"}, {"field": "campaign_name"}, {"field": "sent", "aggregate": "sum"}],
                                                   "view": {"type": "hierarchy"}})
    kept = definition.validate(DATASET, drill.drilled(DATASET, stacked_levels, [{"field": "submitted_at", "bucket": "year", "value": "2026-01-01"}], keep_rest=True))
    check("a hierarchy drilled into 2026 shows its quarters, then channel and campaign under them",
          [(c["field"], c.get("bucket")) for c in kept["columns"]] == [("submitted_at", "quarter"), ("channel", None), ("campaign_name", None), ("sent", None)], str(kept["columns"]))
    by_channel = definition.validate(DATASET, drill.drilled(DATASET, stacked_levels, [{"field": "submitted_at", "bucket": "year", "value": "2026-01-01"}, {"field": "submitted_at", "bucket": "quarter", "value": "2026-07-01"},
                                                                                     {"field": "submitted_at", "bucket": "month", "value": "2026-09-01"}, {"field": "submitted_at", "bucket": "day", "value": "2026-09-02"}], keep_rest=True))
    check("past the days, channel with campaign under it", [c["field"] for c in by_channel["columns"]] == ["channel", "campaign_name", "sent"])

    # range slicers: a number or a date between two ends, or from one, as rules where the dataset has the field
    check("a range comes from the body as its ends; an empty end is left out", dashboards._slicers_body({"sent": {"from": 100, "to": ""}, "channel": ["SMS"]}) == {"sent": {"from": 100}, "channel": ["SMS"]})
    check("ends become between, at least or at most", dashboards._slicer_rules(DATASET, {"sent": {"from": 100, "to": 500}, "submitted_at": {"to": "2026-09-30"}, "messages": {"from": 2}})
          == [{"field": "sent", "operator": "between", "value": [100, 500]}, {"field": "submitted_at", "operator": "<=", "value": "2026-09-30"}, {"field": "messages", "operator": ">=", "value": 2}])
    check("a range slicer is a number field or a date dimension; a text field is not",
          dashboards._slicers_layout([{"field": "sent", "kind": "range"}, {"field": "submitted_at", "kind": "range"}], [DATASET])
          == [{"field": "sent", "label": "Sent Count", "kind": "range", "type": "number"}, {"field": "submitted_at", "label": "Sent", "kind": "range", "type": "datetime"}])
    try:
        dashboards._slicers_layout([{"field": "channel", "kind": "range"}], [DATASET])
        check("a range on text is refused", False)
    except Exception as exc:                                                       # noqa: BLE001
        check("a range on text is refused", "layout.slicers[0]" in str(exc))
    check("the tiles of every page, each with its page's title", dashboards._tiles({"rows": [], "pages": [{"id": "p1", "title": "Overview", "rows": [{"width": "1", "tiles": [{"id": "t1"}]}]},
                                                                                                        {"id": "p2", "title": "Trends", "rows": [{"width": "1", "tiles": [{"id": "t2"}]}]}]})
          == [{"id": "t1", "_page": "Overview"}, {"id": "t2", "_page": "Trends"}])

    # linked datasets: the catalogue's link, the blend checked, the other side's definition, the rows matched
    orders = {"key": "attributed_orders", "label": "Attributed Orders", "fields": [
        {"key": "campaign_name", "label": "Credited To", "type": "string", "role": "dimension", "aggregates": [], "format": None, "sensitivity": "NORMAL"},
        {"key": "channel", "label": "Channel", "type": "string", "role": "dimension", "aggregates": [], "format": None, "sensitivity": "NORMAL"},
        {"key": "ordered_at", "label": "Placed", "type": "datetime", "role": "dimension", "aggregates": [], "format": "datetime", "sensitivity": "NORMAL"},
        {"key": "amount", "label": "Amount", "type": "number", "role": "measure", "aggregates": ["sum", "avg"], "format": "money", "sensitivity": "NORMAL"},
        {"key": "buyers", "label": "Buyers", "type": "number", "role": "measure", "aggregates": ["count"], "format": None, "sensitivity": "PII"}]}
    link = {"from": "messages", "to": "attributed_orders", "label": "Orders credited", "pairs": [["campaign_name", "campaign_name"], ["channel", "channel"]], "time_pair": ["submitted_at", "ordered_at"]}
    attached = links.attach([DATASET, orders], [link])
    linked = attached[0]
    check("the catalogue gives a dataset its links with the other side's measures, personal ones left out",
          linked["links"][0]["to"] == "attributed_orders" and [m["key"] for m in linked["links"][0]["measures"]] == ["amount"] and "links" not in attached[1])
    blended = definition.validate(linked, {"columns": [{"field": "campaign_name"}, {"field": "submitted_at", "bucket": "month"}, {"field": "sent", "aggregate": "sum"}],
                                           "filters": {"combinator": "and", "rules": [{"field": "channel", "operator": "=", "value": "EMAIL"}, {"field": "submitted_at", "operator": "last_days", "value": 30},
                                                                                      {"field": "party_id", "operator": "is_not_null"}]},
                                           "orderBy": [{"field": "attributed_orders.amount", "direction": "desc"}],
                                           "blend": {"dataset": "attributed_orders", "columns": [{"field": "amount", "aggregate": "sum"}]}})
    check("a blend keeps its measures, and the report may sort by one of them",
          blended["blend"] == {"dataset": "attributed_orders", "columns": [{"field": "amount", "aggregate": "sum"}]} and blended["orderBy"] == [{"field": "attributed_orders.amount", "direction": "desc"}])
    other_def, own = links.other_definition(linked, blended)
    check("the other side groups by the same things renamed, takes the rules on what both share, and leaves the rest to the report's side",
          other_def["columns"] == [{"field": "campaign_name"}, {"field": "ordered_at", "bucket": "month"}, {"field": "amount", "aggregate": "sum"}]
          and other_def["filters"]["rules"] == [{"field": "channel", "operator": "=", "value": "EMAIL"}, {"field": "ordered_at", "operator": "last_days", "value": 30}] and own == ["party_id"], str(other_def))
    check("the report's own query leaves the sort on the linked measure for later", 'ORDER BY "c0" ASC NULLS LAST' in compiler.compile(linked, blended, T, "America/New_York", limit=10).sql)
    merged = links.merge(blended, [["A", date(2026, 9, 1), 10], ["B", date(2026, 9, 1), 5]], [["A", date(2026, 9, 1), Decimal("99.5")], ["C", date(2026, 9, 1), 7]])
    check("each row gets the other side's total for its groups, empty where there is none; the other side's rows with no match are not shown",
          merged == [["A", date(2026, 9, 1), 10, Decimal("99.5")], ["B", date(2026, 9, 1), 5, None]], str(merged))
    check("the blended column is named after its dataset", links.columns(linked, blended, orders)[0] | {} == {"key": "attributed_orders.amount", "label": "Amount (Attributed Orders)", "type": "number",
                                                                                                              "format": "money", "aggregate": "sum", "linked": "attributed_orders"})

    def refused_linked(raw, where):
        try:
            definition.validate(linked, raw)
            return False
        except DefinitionError as exc:
            return exc.field == where
    plus = {"dataset": "attributed_orders", "columns": [{"field": "amount", "aggregate": "sum"}]}
    check("a blend groups only by what both share, a date only by a period",
          refused_linked({"columns": [{"field": "party_id"}, {"field": "sent", "aggregate": "sum"}], "blend": plus}, "columns[0].field")
          and refused_linked({"columns": [{"field": "submitted_at"}, {"field": "sent", "aggregate": "sum"}], "blend": plus}, "columns[0].field"))
    check("a blend has no comparisons, no Top N, no hierarchy, and only the other side's offered measures",
          refused_linked({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum", "compare": ["percent_of_total"]}], "blend": plus}, "blend")
          and refused_linked({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "top": {"n": 2}, "blend": plus}, "blend")
          and refused_linked({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "hierarchy"}, "blend": plus}, "blend")
          and refused_linked({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "blend": {"dataset": "attributed_orders", "columns": [{"field": "buyers", "aggregate": "count"}]}}, "blend.columns[0].field")
          and refused_linked({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "blend": {"dataset": "contacts", "columns": []}}, "blend.dataset"))
    check("a drill carries the blend down", definition.validate(linked, drill.drilled(linked, definition.validate(linked, {"columns": [{"field": "submitted_at", "bucket": "year"}, {"field": "sent", "aggregate": "sum"}], "blend": plus}),
                                                                                    [{"field": "submitted_at", "bucket": "year", "value": "2026-01-01"}]))["blend"] == plus)

    # Show the SQL: the statement with its values written in; the escaped percent sign of the statement, not the value's
    readable = ai_report.readable(c)
    check("Show the SQL writes every value in", "%s" not in readable and "'EMAIL'" in readable and "('SMS', 'EMAIL')" in readable and readable.endswith("LIMIT 51 OFFSET 0")
          and "'%Fall \\%%'" in readable, readable)
    check("a ClickHouse statement's named values are written in, quotes doubled",
          ai_report.readable(compiler.Compiled("SELECT {p0:String}, {p1:Int64}", {"p0": "O'Neil", "p1": 5}, [])) == "SELECT 'O''Neil', 5")
    changed = ai_report._changed({"kind": "draft", "definition": {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}]}, "view": {"type": "chart", "chart": "bar"}},
                                 DATASET, {"top": {"n": 3, "others": True}, "view": {"type": "grid"}}, 1000)
    check("a change keeps what the model cannot touch (Top N here) and takes the new view", changed["definition"]["top"] == {"n": 3, "others": True}
          and changed["view"]["chart"] == "bar")
    table, left_out = ai_report.rows_text({"columns": [{"label": "Campaign"}, {"label": "Customer", "masked": True}, {"label": "Sent"}], "rows": [["Fall", "•••", 12.0], ["Spring", "•••", 1234.5]]})
    check("explaining sends the rows without the masked columns", table == "Campaign | Sent\nFall | 12\nSpring | 1,234.5" and left_out == ["Customer"], table)
    check("an embed keeps only the secret's hash", len(embeds.digest("abc")) == 64 and embeds.digest("abc") == embeds.digest("abc") != embeds.digest("abd"))
    from django.test import RequestFactory
    from apps.tenants import resolver
    with mock.patch.object(resolver, "_lookup", side_effect=lambda key, value: (key, value)):
        check("an embed's address names its workspace", resolver.resolve_tenant(RequestFactory().get("/api/v1/public/embed/NWA/s3cret/")) == ("code", "NWA"))

    # the visuals: chart kinds and their settings saved with the view, the engine's grand total for a KPI card or a gauge
    kpi = definition.validate(DATASET, {"columns": [{"field": "submitted_at", "bucket": "month"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "chart": "kpi", "target": 900}})
    check("a KPI card keeps its target, and the engine totals the whole period for it", kpi["view"] == {"type": "chart", "chart": "kpi", "target": 900}
          and compiler.totals_sets(kpi, kpi["view"]) == [()])
    many = definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "submitted_at", "bucket": "month"}, {"field": "sent", "aggregate": "sum"}, {"field": "messages", "aggregate": "count"}],
                                         "view": {"type": "pivot", "across": "submitted_at", "matrix": True, "sparklines": True}})
    check("a pivot keeps All Values and the Trend column", many["view"] == {"type": "pivot", "across": "submitted_at", "matrix": True, "sparklines": True})
    check("small multiples repeat by a dimension of the report; a map keeps its geography; an analysing view keeps its value",
          definition.validate(DATASET, {**{k: many[k] for k in ("columns",)}, "view": {"type": "chart", "chart": "bar", "multiples": "channel"}})["view"]["multiples"] == "channel"
          and definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "chart": "map", "geo": "world"}})["view"]["geo"] == "world"
          and definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "influencers", "measure": "sent"}})["view"] == {"type": "influencers", "measure": "sent"})
    check("a wrong setting is refused where it sits",
          refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "chart": "gauge", "target": "high"}}, "view.target")
          and refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "chart": "bar", "multiples": "sent"}}, "view.multiples")
          and refused({"columns": [{"field": "channel"}, {"field": "campaign_name"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "chart": "funnel", "multiples": "campaign_name"}}, "view.multiples")
          and refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "chart": "map", "geo": "mars"}}, "view.geo")
          and refused({"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "decomposition", "measure": "channel"}}, "view.measure")
          and refused({"columns": [{"field": "channel"}], "view": {"type": "chart", "chart": "radar"}}, "view.chart"))
    explained_by = analyze.candidates(DATASET, definition.validate(DATASET, {"columns": [{"field": "campaign_name"}, {"field": "sent", "aggregate": "sum"}],
                                                                            "filters": {"combinator": "and", "rules": [{"field": "channel", "operator": "=", "value": "EMAIL"}]}}), mask=True)
    check("Key Influencers explains by the report's dimensions first, a date by month, never a field fixed to one value, the dataset's key or personal data hidden from the person",
          explained_by == [{"field": "campaign_name", "label": "Campaign"}, {"field": "submitted_at", "label": "Sent", "bucket": "month"}], str(explained_by))
    columns_k = [{"key": "submitted_at", "label": "Sent", "type": "date", "format": "date", "bucket": "month"}, {"key": "sent", "label": "Sent Count", "type": "number", "format": "integer", "aggregate": "sum"}]
    figure_pdf = pdf.report("KPI", columns_k, [["2026-09-01", 400], ["2026-10-01", 500]], {"type": "chart", "chart": "kpi", "target": 1000}, [{"keys": {}, "values": [900]}], workspace="W", when="now")
    water_pdf = pdf.report("Steps", columns_k, [["2026-09-01", 400], ["2026-10-01", -50]], {"type": "chart", "chart": "waterfall"}, workspace="W", when="now")
    scatter_pdf = pdf.report("Points", columns_k, [["2026-09-01", 400]], {"type": "chart", "chart": "scatter"}, workspace="W", when="now")
    summary_pdf = pdf.dashboard("Board", None, "September", [{"title": "Summary", "kind": "summary", "summary": "Sends rose <fast> & steadily.", "points": ["Email led."]}], workspace="W", when="now")
    check("the PDF draws a KPI card, a waterfall as steps, the rest as rows, and a summary tile (its words escaped)",
          all(b[:5] == b"%PDF-" for b in (figure_pdf, water_pdf, scatter_pdf, summary_pdf)))

    # dashboards and interaction: a tooltip report on a view, alerts' arithmetic, a comment's anchor, PowerPoint and the print-ready PDF
    check("a view may name a tooltip report", definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}],
                                                                           "view": {"type": "chart", "tooltip": "A0000000-0000-4000-8000-000000000009"}})["view"]["tooltip"] == "a0000000-0000-4000-8000-000000000009"
          and refused({"columns": [{"field": "channel"}], "view": {"type": "chart", "tooltip": "rep-1"}}, "view.tooltip"))
    from datetime import datetime as _dt, time as _time, timezone as _tz
    from zoneinfo import ZoneInfo
    check("an alert's condition, and when it is checked next", alerts._met(87.2, "<", 90) is True and alerts._met(90, "<", 90) is False and alerts._met(90, "<=", 90) is True
          and alerts._met(None, ">", 1) is None
          and alerts.next_check("hour", None, ZoneInfo("UTC"), _dt(2026, 10, 19, 9, 20, 30, tzinfo=_tz.utc)) == _dt(2026, 10, 19, 10, 20, tzinfo=_tz.utc)
          and alerts.next_check("day", _time(8, 0), ZoneInfo("America/New_York"), _dt(2026, 10, 19, 13, 0, tzinfo=_tz.utc)) == _dt(2026, 10, 20, 12, 0, tzinfo=_tz.utc))
    check("a comment's anchor: a tile, or a field's value with its words; the whole item when empty",
          comments._anchor({"tile_id": "t1"}) == {"tile_id": "t1"} and comments._anchor({"field": "channel", "value": "SMS"}) == {"field": "channel", "value": "SMS", "label": "SMS"}
          and comments._anchor({}) is None)
    import io as _io
    from pptx import Presentation
    columns_s = [{"key": "channel", "label": "Channel", "type": "string", "format": None}, {"key": "sent", "label": "Sent Count", "type": "number", "format": "integer", "aggregate": "sum"}]
    deck = Presentation(_io.BytesIO(slides.report("Sends", columns_s, [["EMAIL", 900], ["SMS", 40]], {"type": "chart", "chart": "bar"}, workspace="W", when="now")))
    check("a report as PowerPoint: a cover, its chart as a native chart, then its rows", len(deck.slides) == 3 and any(sh.has_chart for sh in deck.slides[1].shapes)
          and any(sh.has_table for sh in deck.slides[2].shapes))
    board_deck = Presentation(_io.BytesIO(slides.dashboard("Board", None, "September", [
        {"title": "Sent", "kind": "figure", "value": 940, "previous": 800, "compared": True, "column": columns_s[1]},
        {"title": "By channel", "kind": "view", "view": {"type": "chart", "chart": "pie"}, "result": {"columns": columns_s, "rows": [["EMAIL", 900], ["SMS", 40]], "row_count": 2}},
        {"title": "Note", "kind": "text", "text": "Read me first."}, {"title": "Gone", "kind": "missing", "detail": "The report is gone."}], workspace="W", when="now")))
    check("a dashboard as PowerPoint: a cover and a slide per tile", len(board_deck.slides) == 5)
    many = [[f"C{i % 3}", i] for i in range(400)]
    printed = pdf.paginated("Long", columns_s, many, None, size="letter", wide=True, group=True, filters=["Channel is EMAIL"], workspace="W", when="now")
    import re as _re
    plain = pdf.paginated("Long", columns_s, many, None, workspace="W", when="now")
    check("the print-ready PDF lays every row over pages; a section per value adds pages", printed[:5] == b"%PDF-"
          and len(_re.findall(rb"/Type /Page[^s]", plain)) >= 8 and len(_re.findall(rb"/Type /Page[^s]", printed)) > len(_re.findall(rb"/Type /Page[^s]", plain)))

    # drill-down: the levels, and the definition of the level under a path
    yearly = definition.validate(DATASET, {"columns": [{"field": "submitted_at", "bucket": "year"}, {"field": "channel"}, {"field": "campaign_name"}, {"field": "sent", "aggregate": "sum", "compare": ["percent_of_total"]}],
                                           "filters": {"combinator": "and", "rules": [{"field": "channel", "operator": "!=", "value": "PUSH"}]}, "orderBy": [{"field": "sent", "direction": "desc"}],
                                           "view": {"type": "chart", "chart": "bar", "across": "channel"}})
    found = drill.levels(DATASET, yearly)
    check("a year drills to quarters, months and days, then the next dimension; the one spread across stays the series",
          [(l["field"], l.get("bucket")) for l in found] == [("submitted_at", "year"), ("submitted_at", "quarter"), ("submitted_at", "month"), ("submitted_at", "day"), ("campaign_name", None)]
          and found[1]["label"] == "Sent by Quarter")
    below = definition.validate(DATASET, drill.drilled(DATASET, yearly, [{"field": "submitted_at", "bucket": "year", "value": "2026-01-01"}, {"field": "submitted_at", "bucket": "quarter", "value": "2026-07-01"}]))
    check("under 2026 and Q3: months as the rows, the series kept, the measure kept, the quarter's dates as a rule on top of the report's own",
          [(c["field"], c.get("bucket")) for c in below["columns"]] == [("submitted_at", "month"), ("channel", None), ("sent", None)]
          and below["filters"]["rules"][-2:] == [{"field": "submitted_at", "operator": "between", "value": ["2026-01-01", "2026-12-31"]}, {"field": "submitted_at", "operator": "between", "value": ["2026-07-01", "2026-09-30"]}]
          and below["filters"]["rules"][0]["field"] == "channel" and below["orderBy"] == [{"field": "submitted_at", "direction": "asc"}] and below["view"]["across"] == "channel", str(below))
    by_campaign = definition.validate(DATASET, drill.drilled(DATASET, yearly, [{"field": "submitted_at", "bucket": "year", "value": "2026-01-01"}, {"field": "submitted_at", "bucket": "quarter", "value": "2026-07-01"},
                                                                              {"field": "submitted_at", "bucket": "month", "value": "2026-09-01"}, {"field": "submitted_at", "bucket": "day", "value": "2026-09-15"}]))
    check("past the days, the next dimension, sorted as the report was; a share of the total survives the drill", by_campaign["columns"][0]["field"] == "campaign_name"
          and by_campaign["orderBy"] == [{"field": "sent", "direction": "desc"}] and by_campaign["columns"][2]["compare"] == ["percent_of_total"])
    for path, why in (([{"field": "channel", "value": "EMAIL"}], "a step that is not the level"), ([{"field": "submitted_at", "bucket": "year", "value": "x"}], "a date step that is not a date")):
        try:
            drill.drilled(DATASET, yearly, path)
            check(f"{why} is refused", False)
        except drill.DrillError:
            check(f"{why} is refused", True)
    check("one level is nothing to drill", drill.info(DATASET, definition.validate(DATASET, {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}]}), []) is None
          and drill.bucket_end(date(2026, 7, 1), "quarter") == date(2026, 9, 30) and drill.bucket_end(date(2026, 2, 1), "month") == date(2026, 2, 28))
    check.done()


def _tables_present() -> bool:
    with connections[P].cursor() as cursor:
        cursor.execute("SELECT count(*) FROM information_schema.tables WHERE table_schema = 'public' AND table_name IN ('report_dataset', 'report', 'report_run_audit', 'dashboard', 'report_favourite', 'report_schedule')")
        return cursor.fetchone()[0] == 6


def test_api(check, world):
    if not _tables_present():
        pytest.skip("patches 0049 to 0052 have not been applied to this database")
    me = world.api("get", "/api/v1/me/").json()
    check("the workspace has Insights (Beta, on every plan) and /me says how it is offered", any(c["key"] == "insights" for c in me["components"]) and me["insights"]["mode"] in ("integrated", "standalone", "both"), str(me.get("insights")))

    cat = world.api("get", "/api/v1/insights/catalog/")
    body = cat.json() if cat.status_code == 200 else {}
    keys = [d["key"] for d in body.get("datasets", []) if "kind" not in d]          # the platform's; the workspace's own (uploads, its databases) carry a kind
    check("the catalogue lists the first datasets with their fields and the engine", cat.status_code == 200 and keys == ["messages", "attributed_orders", "contacts", "journey_enrollments"]
          and body["engine"]["engine"] == "postgresql" and body["engine"]["limits"]["preview_rows"] > 0, f"{cat.status_code} {keys}")
    messages = next(d for d in body["datasets"] if d["key"] == "messages")
    check("an administrator sees PII fields; the fields carry their operators and buckets", any(f["key"] == "party_id" for f in messages["fields"])
          and next(f for f in messages["fields"] if f["key"] == "submitted_at")["buckets"][:2] == ["day", "week"])

    draft = {"dataset": "messages", "definition": {"columns": [{"field": "campaign_name"}, {"field": "channel"}, {"field": "sent", "aggregate": "sum"}, {"field": "messages", "aggregate": "count"}],
                                                   "filters": {"combinator": "and", "rules": [{"field": "test_mode", "operator": "=", "value": False}]},
                                                   "orderBy": [{"field": "messages", "direction": "desc"}]}}
    preview = world.api("post", "/api/v1/insights/reports/preview/", draft)
    rows = preview.json() if preview.status_code == 200 else {}
    check("a preview runs the draft under the preview cap and answers columns and rows", preview.status_code == 200 and [c["label"] for c in rows.get("columns", [])] == ["Campaign", "Channel", "Sent Count", "Messages"]
          and isinstance(rows.get("rows"), list) and rows["limit"] <= body["engine"]["limits"]["preview_rows"] and rows["cached"] is False, preview.content[:300])
    again = world.api("post", "/api/v1/insights/reports/preview/", draft).json()
    check("the same draft again comes from the cache", again.get("cached") is True)
    bad = world.api("post", "/api/v1/insights/reports/preview/", {"dataset": "messages", "definition": {"columns": [{"field": "sent"}]}})
    check("a draft the catalogue refuses is a 400 naming the part", bad.status_code == 400 and "definition.columns[0].aggregate" in json.dumps(bad.json()), bad.content[:200])

    created = world.api("post", "/api/v1/insights/reports/", {"name": "Sends by campaign", **draft})
    report = created.json() if created.status_code == 201 else {}
    check("a report is created as version 1, private, owned by its author", created.status_code == 201 and report["version"] == 1 and report["visibility"] == "PRIVATE" and report["owner_name"], created.content[:200])
    rid = report.get("id")
    listed = world.api("get", "/api/v1/insights/reports/?component=cep").json()["reports"]
    check("the list shows it under its component with its dataset's label", any(r["id"] == rid and r["component"]["key"] == "cep" and r["dataset_label"] == "Messages" and r["can_edit"] for r in listed))
    run = world.api("post", f"/api/v1/insights/reports/{rid}/run/", {"limit": 10})
    check("a saved report runs and is stamped", run.status_code == 200 and run.json()["limit"] == 10 and world.api("get", f"/api/v1/insights/reports/{rid}/").json()["last_run_at"], run.content[:200])
    as_pivot = world.api("post", f"/api/v1/insights/reports/{rid}/run/", {"limit": 10, "view": {"type": "pivot", "across": "channel"}})
    pv_body = as_pivot.json() if as_pivot.status_code == 200 else {}
    check("run as a pivot, the engine totals the rows, the columns and the whole", as_pivot.status_code == 200 and isinstance(pv_body.get("totals"), list)
          and any(t["keys"] == {} and len(t["values"]) == 2 for t in pv_body["totals"]) and all(set(t["keys"]) <= {"campaign_name", "channel"} for t in pv_body["totals"]), as_pivot.content[:300])
    # time comparisons and drill-down through the API: a monthly report compared with the month before and last year, then drilled
    monthly = world.api("post", "/api/v1/insights/reports/", {"name": "Sends by year and channel", "dataset": "messages", "definition": {
        "columns": [{"field": "submitted_at", "bucket": "year"}, {"field": "channel"}, {"field": "sent", "aggregate": "sum", "compare": ["previous_period", "change_previous_year", "percent_of_total"]}],
        "filters": {"combinator": "and", "rules": [{"field": "submitted_at", "operator": "last_days", "value": 400}]}, "orderBy": [{"field": "submitted_at", "direction": "asc"}], "view": {"type": "chart", "chart": "bar"}}})
    mid = monthly.json().get("id") if monthly.status_code == 201 else None
    compared = world.api("post", f"/api/v1/insights/reports/{mid}/run/", {"limit": 50})
    cmp_body = compared.json() if compared.status_code == 200 else {}
    check("a compared report runs: each comparison a column after its measure, the drill levels beside the rows", compared.status_code == 200
          and [c["key"] for c in cmp_body["columns"]] == ["submitted_at", "channel", "sent", "sent__previous_period", "sent__change_previous_year", "sent__percent_of_total"]
          and all(len(r) == 6 for r in cmp_body["rows"]) and [l["label"] for l in cmp_body["drill"]["levels"]][:2] == ["Sent by Year", "Sent by Quarter"] and cmp_body["drill"]["depth"] == 0, compared.content[:400])
    first = (cmp_body.get("rows") or [["2026-01-01"]])[0][0] or "2026-01-01"
    drilled_run = world.api("post", f"/api/v1/insights/reports/{mid}/run/", {"limit": 50, "drill": {"path": [{"field": "submitted_at", "bucket": "year", "value": first}]}})
    dr_body = drilled_run.json() if drilled_run.status_code == 200 else {}
    check("drilled into a year, the rows are its quarters and the comparisons are against the quarter before", drilled_run.status_code == 200
          and [c["key"] for c in dr_body["columns"]][:3] == ["submitted_at", "sent", "sent__previous_period"] and dr_body["columns"][0].get("bucket") == "quarter" and dr_body["columns"][2]["label"] == "Sent Count, Previous Quarter" and dr_body["drill"]["depth"] == 1
          and all(str(r[0]) >= str(first)[:4] for r in dr_body["rows"]), drilled_run.content[:400])
    wrong = world.api("post", f"/api/v1/insights/reports/{mid}/run/", {"drill": {"path": [{"field": "channel", "value": "EMAIL"}]}})
    check("a drill path that does not fit the report is refused under drill", wrong.status_code == 400 and "drill" in json.dumps(wrong.json()), wrong.content[:200])
    csv_cmp = world.api("get", f"/api/v1/insights/reports/{mid}/export.csv")
    check("the CSV export carries the comparison columns", csv_cmp.status_code == 200 and "Sent Count, Change vs Last Year" in b"".join(csv_cmp.streaming_content).decode(), str(csv_cmp.status_code))
    # drill-through through the API: a detail report taking Channel, listed for the field, opened filtered to a value
    listed_before = {r["id"] for r in world.api("get", "/api/v1/insights/reports/?dataset=messages&detail_for=channel").json()["reports"]}
    detail_r = world.api("post", "/api/v1/insights/reports/", {"name": "Campaigns for a channel", "dataset": "messages", "definition": {
        "columns": [{"field": "campaign_name"}, {"field": "sent", "aggregate": "sum"}], "drillthrough": ["channel"]}})
    did_r = detail_r.json().get("id") if detail_r.status_code == 201 else None
    listed_dt = world.api("get", "/api/v1/insights/reports/?dataset=messages&detail_for=channel").json()["reports"]
    check("the detail reports for a field are listed for it, and only those", detail_r.status_code == 201 and {r["id"] for r in listed_dt} == listed_before | {did_r}
          and mid not in listed_before, str([r["name"] for r in listed_dt]))
    through = world.api("post", f"/api/v1/insights/reports/{did_r}/run/", {"through": [{"field": "channel", "value": "EMAIL"}]})
    check("opened from another report, the detail report runs filtered to the value and says so", through.status_code == 200
          and through.json()["through"] == [{"field": "channel", "label": "Channel", "value": "EMAIL"}], through.content[:300])
    refused_through = world.api("post", f"/api/v1/insights/reports/{did_r}/run/", {"through": [{"field": "campaign_name", "value": "x"}]})
    check("a field the detail report does not take is refused", refused_through.status_code == 400 and "through[0]" in json.dumps(refused_through.json()), refused_through.content[:200])

    # sharing and row access (patch 0060): only where its tables are there
    if access.present("report_row_rule") and access.present("dashboard_share"):
        people = world.api("get", "/api/v1/insights/principals/?q=").json()
        me_id = me["user"]["id"] if "user" in me else me.get("principal", {}).get("user_id")
        check("people and roles to share with are listed", isinstance(people.get("people"), list) and isinstance(people.get("roles"), list), str(people)[:200])
        role_id = people["roles"][0]["id"] if people["roles"] else None
        if role_id:
            put = world.api("put", f"/api/v1/insights/reports/{did_r}/shares/", {"visibility": "PRIVATE", "shares": [{"kind": "role", "id": role_id, "permission": "edit"}]})
            check("a report is shared with a role to edit, and kept private", put.status_code == 200 and put.json()["visibility"] == "PRIVATE"
                  and [(s["kind"], s["permission"]) for s in put.json()["shares"]] == [("role", "edit")], put.content[:300])
            check("the list marks it shared", any(r["id"] == did_r and r["shared"] for r in world.api("get", "/api/v1/insights/reports/").json()["reports"]))
            bad_share = world.api("put", f"/api/v1/insights/reports/{did_r}/shares/", {"shares": [{"kind": "role", "id": "00000000-0000-4000-8000-000000000000"}]})
            check("a role of no one is refused where it sits", bad_share.status_code == 400 and "shares[0]" in json.dumps(bad_share.json()))
        rules = world.api("get", "/api/v1/insights/row-rules/").json()
        choices = next(d for d in rules["datasets"] if d["key"] == "messages")
        check("the datasets offer their text, number and yes/no dimensions for a rule", any(f["key"] == "channel" for f in choices["fields"]) and not any(f["key"] == "sent" for f in choices["fields"]))
        user_id = people["people"][0]["id"] if people["people"] else None
        made_rule = world.api("post", "/api/v1/insights/row-rules/", {"dataset": "messages", "field": "channel", "principal": {"kind": "user", "id": user_id}, "values": ["SMS"]})
        check("a rule limits a person to SMS", made_rule.status_code == 201 and made_rule.json()["values"] == ["SMS"], made_rule.content[:300])
        twice = world.api("post", "/api/v1/insights/row-rules/", {"dataset": "messages", "field": "channel", "principal": {"kind": "user", "id": user_id}, "values": ["EMAIL"]})
        check("a second rule for the same person and field is refused", twice.status_code == 400)
        if made_rule.status_code == 201:
            check("the rule goes", world.api("delete", f"/api/v1/insights/row-rules/{made_rule.json()['id']}/").status_code == 204)
    changed = world.api("put", f"/api/v1/insights/reports/{rid}/", {"row_version": report["row_version"], "name": "Sends by campaign and channel", "visibility": "WORKSPACE",
                                                                   "definition": {**draft["definition"], "limit": 100, "view": {"type": "pivot", "across": "channel"}}, "note": "Shared with the workspace"})
    check("the report's view follows the definition", changed.status_code == 200 and changed.json()["view_type"] == "pivot" and changed.json()["definition"]["view"] == {"type": "pivot", "across": "channel"}, changed.content[:200])
    check("a change makes version 2 and keeps the history", changed.status_code == 200 and changed.json()["version"] == 2 and changed.json()["visibility"] == "WORKSPACE"
          and [v["version"] for v in world.api("get", f"/api/v1/insights/reports/{rid}/versions/").json()["versions"]] == [2, 1], changed.content[:200])
    stale = world.api("put", f"/api/v1/insights/reports/{rid}/", {"row_version": report["row_version"], "name": "Too late"})
    check("a stale change is a 409", stale.status_code == 409)
    history = world.api("get", f"/api/v1/insights/reports/{rid}/versions/").json()["versions"]
    check("every version carries its definition", history[0]["version"] == 2 and history[0]["definition"]["limit"] == 100 and "limit" not in history[1]["definition"], str(history)[:200])
    restored = world.api("post", f"/api/v1/insights/reports/{rid}/versions/1/restore/", {})
    check("restoring version 1 makes version 3 from its definition, and says so", restored.status_code == 200 and restored.json()["version"] == 3 and "limit" not in restored.json()["definition"]
          and world.api("get", f"/api/v1/insights/reports/{rid}/versions/").json()["versions"][0]["note"] == "Restored from version 1", restored.content[:200])
    check("a version that never was is 404", world.api("post", f"/api/v1/insights/reports/{rid}/versions/9/restore/", {}).status_code == 404)
    export = world.api("get", f"/api/v1/insights/reports/{rid}/export.csv")
    text = b"".join(export.streaming_content).decode() if export.status_code == 200 else ""
    check("the export streams a CSV with the labels as its header", export.status_code == 200 and text.startswith("Campaign,Channel,Sent Count,Messages") and "attachment" in export["Content-Disposition"], text[:120])
    paper = world.api("get", f"/api/v1/insights/reports/{rid}/export.pdf")
    check("the PDF export is a PDF named after the report", paper.status_code == 200 and paper["Content-Type"] == "application/pdf" and paper.content[:5] == b"%PDF-", str(paper.status_code))
    book = world.api("get", f"/api/v1/insights/reports/{rid}/export.xlsx")
    check("the Excel export is a workbook named after the report", book.status_code == 200 and book["Content-Type"].startswith("application/vnd.openxmlformats") and 'filename="Sends by campaign and channel.xlsx"' in book["Content-Disposition"]
          and book.content[:2] == b"PK", str(book.status_code))

    # scheduled delivery (Phase E): a schedule on the report, a test sending through the (console) mail backend, pause and resume
    made_s = world.api("post", "/api/v1/insights/schedules/", {"item_id": rid, "every": "week", "on_day": 1, "at_time": "08:00", "format": "xlsx", "recipients": ["finance@northwind.example", "Aby@Northwind.example"],
                                                                "subject": "Sends · {date}", "note": "Last week's sends."})
    sched = made_s.json() if made_s.status_code == 201 else {}
    check("a schedule is made with its next sending computed and its addresses cleaned", made_s.status_code == 201 and sched["words"] == "Every Monday at 08:00" and sched["status"] == "ACTIVE"
          and sched["recipients"] == ["finance@northwind.example", "aby@northwind.example"] and sched["next_run_at"], made_s.content[:300])
    sid = sched.get("id")
    check("a bad address, and an unknown format, are refused where they sit",
          "recipients" in json.dumps(world.api("post", "/api/v1/insights/schedules/", {"item_id": rid, "every": "day", "at_time": "08:00", "recipients": ["nope"]}).json())
          and "format" in json.dumps(world.api("post", "/api/v1/insights/schedules/", {"item_id": rid, "every": "day", "at_time": "08:00", "format": "docx", "recipients": ["a@b.co"]}).json()))
    test_send = world.api("post", f"/api/v1/insights/schedules/{sid}/test/", {})
    check("Send a Test Now builds the file and sends it to the asker, recorded as a test delivery and audited as a delivery", test_send.status_code == 200 and test_send.json()["outcome"] == "OK"
          and test_send.json()["recipients"] == 1 and world.sql("SELECT trigger, outcome FROM report_delivery WHERE schedule_id = %s", [sid], fetch=True) == [("test", "OK")]
          and ("delivery", "OK") in world.sql("SELECT kind, outcome FROM report_run_audit WHERE report_id = %s", [rid], fetch=True), test_send.content[:300])
    paused = world.api("post", f"/api/v1/insights/schedules/{sid}/pause/", {})
    resumed = world.api("post", f"/api/v1/insights/schedules/{sid}/resume/", {})
    check("pause clears the next sending; resume computes it again", paused.json()["status"] == "PAUSED" and paused.json()["next_run_at"] is None and resumed.json()["status"] == "ACTIVE" and resumed.json()["next_run_at"])
    listed_s = world.api("get", f"/api/v1/insights/schedules/?item={rid}").json()
    check("the list shows the report's schedule and whether mail is set up", any(s["id"] == sid for s in listed_s["schedules"]) and "mail_configured" in listed_s)
    check("a deleted schedule is gone", world.api("delete", f"/api/v1/insights/schedules/{sid}/").status_code == 204 and not world.api("get", f"/api/v1/insights/schedules/?item={rid}").json()["schedules"])

    # favourites and folders (Phase C)
    made_folder = world.api("post", "/api/v1/insights/folders/", {"name": "Finance"})
    folder = made_folder.json() if made_folder.status_code == 201 else {}
    check("a folder is made for the workspace", made_folder.status_code == 201 and folder["name"] == "Finance" and folder["reports"] == 0, made_folder.content[:200])
    twice = world.api("post", "/api/v1/insights/folders/", {"name": "finance"})
    check("a folder name is unique, whatever the case", twice.status_code == 400)
    filed = world.api("put", f"/api/v1/insights/reports/{rid}/", {"row_version": world.api("get", f"/api/v1/insights/reports/{rid}/").json()["row_version"], "folder_id": folder["id"]})
    check("a report is filed in a folder without a new version", filed.status_code == 200 and filed.json()["folder_id"] == folder["id"] and filed.json()["version"] == 3, filed.content[:200])
    check("the star goes on and shows in the list, first", world.api("put", f"/api/v1/insights/favourites/report/{rid}/").status_code == 204
          and next(r for r in world.api("get", "/api/v1/insights/reports/").json()["reports"] if r["id"] == rid)["favourite"] is True)
    in_folder = world.api("get", f"/api/v1/insights/reports/?folder={folder['id']}").json()["reports"]
    none = world.api("get", "/api/v1/insights/reports/?folder=none").json()["reports"]
    check("the list filters by folder, and by no folder", any(r["id"] == rid and r["folder_name"] == "Finance" for r in in_folder) and not any(r["id"] == rid for r in none))
    check("the folders list counts the reports", next(f for f in world.api("get", "/api/v1/insights/folders/").json()["folders"] if f["id"] == folder["id"])["reports"] == 1)
    check("the star comes off", world.api("delete", f"/api/v1/insights/favourites/report/{rid}/").status_code == 204
          and next(r for r in world.api("get", "/api/v1/insights/reports/").json()["reports"] if r["id"] == rid)["favourite"] is False)
    check("deleting the folder leaves its reports, filed nowhere", world.api("delete", f"/api/v1/insights/folders/{folder['id']}/").status_code == 204
          and world.api("get", f"/api/v1/insights/reports/{rid}/").json()["folder_id"] is None)

    # dashboards (Phase C): rows of tiles over the report, one period for every tile
    made = world.api("post", "/api/v1/insights/dashboards/", {"name": "Marketing Weekly"})
    dash = made.json() if made.status_code == 201 else {}
    check("a dashboard starts empty, private, on the last 30 days", made.status_code == 201 and dash["layout"] == {"rows": []} and dash["period"] == {"preset": "30d"}
          and dash["visibility"] == "PRIVATE" and dash["bounds"]["from"] < dash["bounds"]["to"], made.content[:200])
    did = dash.get("id")
    layout = {"rows": [{"width": "2-1", "tiles": [{"id": "t1", "report_id": rid, "show": "chart", "title": "Sends by campaign"}, {"id": "t2", "report_id": rid, "show": "figure", "measure": "sent"}]}]}
    saved = world.api("put", f"/api/v1/insights/dashboards/{did}/", {"row_version": dash["row_version"], "layout": layout, "period": {"preset": "7d"}, "compare": True, "note": "First tiles"})
    check("tiles save as version 2 with the period and the comparison", saved.status_code == 200 and saved.json()["version"] == 2 and saved.json()["compare"] is True
          and len(saved.json()["layout"]["rows"][0]["tiles"]) == 2, saved.content[:300])
    bad = world.api("put", f"/api/v1/insights/dashboards/{did}/", {"row_version": saved.json()["row_version"],
                                                                    "layout": {"rows": [{"width": "1", "tiles": [{"id": "x", "report_id": "00000000-0000-4000-8000-000000000000", "show": "grid"}]}]}})
    check("a tile on a report one cannot open is refused where it sits", bad.status_code == 400 and "layout.rows[0].tiles[0].report_id" in json.dumps(bad.json()), bad.content[:200])
    figure = world.api("post", f"/api/v1/insights/dashboards/{did}/tiles/t2/run/", {})
    fig = figure.json() if figure.status_code == 200 else {}
    check("a figure tile answers the measure over the period, with the previous period beside it", figure.status_code == 200 and fig.get("kind") == "figure" and fig["compared"] is True
          and fig["column"]["key"] == "sent" and fig["own_dates"] is False, figure.content[:300])
    view = world.api("post", f"/api/v1/insights/dashboards/{did}/tiles/t1/run/", {"period": {"preset": "custom", "from": "2026-09-01", "to": "2026-09-30"}})
    vw = view.json() if view.status_code == 200 else {}
    check("a view tile answers the rows in the tile's view for the period asked", view.status_code == 200 and vw.get("kind") == "view" and vw["view"]["type"] == "chart"
          and vw["bounds"] == {"from": "2026-09-01", "to": "2026-09-30", "previous_from": "2026-08-02", "previous_to": "2026-08-31"} and isinstance(vw["result"]["rows"], list), view.content[:300])
    # pages: a dashboard of two pages saves, runs a tile of the second, and exports
    two = world.api("put", f"/api/v1/insights/dashboards/{did}/", {"row_version": world.api("get", f"/api/v1/insights/dashboards/{did}/").json()["row_version"], "layout": {"pages": [
        {"id": "p1", "title": "Overview", "rows": [{"width": "1", "tiles": [{"id": "t1", "report_id": rid, "show": "chart"}]}]},
        {"id": "p2", "title": "Detail", "rows": [{"width": "1", "tiles": [{"id": "t9", "report_id": rid, "show": "grid"}]}]}]}})
    check("a dashboard keeps its pages", two.status_code == 200 and [p["title"] for p in two.json()["layout"]["pages"]] == ["Overview", "Detail"], two.content[:300])
    check("a tile of the second page runs", world.api("post", f"/api/v1/insights/dashboards/{did}/tiles/t9/run/", {}).json().get("kind") == "view")
    dup = world.api("put", f"/api/v1/insights/dashboards/{did}/", {"row_version": two.json()["row_version"], "layout": {"pages": [
        {"id": "p1", "title": "A", "rows": [{"width": "1", "tiles": [{"id": "t1", "report_id": rid, "show": "chart"}]}]}, {"id": "p2", "title": "B", "rows": [{"width": "1", "tiles": [{"id": "t1", "report_id": rid, "show": "grid"}]}]}]}})
    check("a tile id is unique across the pages", dup.status_code == 400, dup.content[:200])
    check("the two-page dashboard exports", world.api("get", f"/api/v1/insights/dashboards/{did}/export.pdf").status_code == 200)
    if access.present("dashboard_view"):
        view = world.api("post", f"/api/v1/insights/dashboards/{did}/views/", {"name": "SMS only", "state": {"selection": {"channel": ["SMS"]}, "page": "p2", "period": {"preset": "7d"}}})
        check("a view saves the slicers, the page and the period", view.status_code == 201 and view.json()["state"]["selection"] == {"channel": ["SMS"]} and view.json()["mine"], view.content[:200])
        check("a person's views are listed", [v["name"] for v in world.api("get", f"/api/v1/insights/dashboards/{did}/views/").json()["views"]] == ["SMS only"])
        check("a view is removed by its owner", world.api("delete", f"/api/v1/insights/dashboards/{did}/views/{view.json()['id']}/").status_code == 204)
    listed_d = world.api("get", "/api/v1/insights/dashboards/").json()["dashboards"]
    check("the list counts the tiles", any(d["id"] == did and d["tiles"] == 2 and d["can_edit"] for d in listed_d), str(listed_d)[:200])
    board_pdf = world.api("get", f"/api/v1/insights/dashboards/{did}/export.pdf")
    check("the dashboard exports as a PDF named after it", board_pdf.status_code == 200 and board_pdf["Content-Type"] == "application/pdf" and board_pdf.content[:5] == b"%PDF-"
          and 'filename="Marketing Weekly.pdf"' in board_pdf["Content-Disposition"], str(board_pdf.status_code))
    dash_sched = world.api("post", "/api/v1/insights/schedules/", {"kind": "dashboard", "item_id": did, "every": "day", "at_time": "07:00", "format": "pdf", "recipients": ["ops@northwind.example"]})
    check("a dashboard is scheduled as PDF, and Send a Test Now sends it", dash_sched.status_code == 201 and dash_sched.json()["item_name"] == "Marketing Weekly"
          and world.api("post", f"/api/v1/insights/schedules/{dash_sched.json()['id']}/test/", {}).json().get("outcome") == "OK", dash_sched.content[:300])
    check("a dashboard as Excel is refused", world.api("post", "/api/v1/insights/schedules/", {"kind": "dashboard", "item_id": did, "every": "day", "at_time": "07:00", "format": "xlsx", "recipients": ["ops@northwind.example"]}).status_code == 400)
    check("a deleted dashboard is gone", world.api("delete", f"/api/v1/insights/dashboards/{did}/").status_code == 204
          and not any(d["id"] == did for d in world.api("get", "/api/v1/insights/dashboards/").json()["dashboards"]))

    audit = world.sql("SELECT kind, outcome FROM report_run_audit WHERE report_id = %s OR (report_id IS NULL AND dataset_key = 'messages') ORDER BY id", [rid], fetch=True)
    check("every preview, run and export is audited with its outcome", ("preview", "OK") in audit and ("run", "OK") in audit and ("export", "OK") in audit, str(audit))
    log = world.api("get", "/api/v1/insights/runs/?days=7&q=Sends by campaign")
    body_log = log.json() if log.status_code == 200 else {}
    check("the run log lists the runs of this report with who, the outcome, the rows and the time, and sums them up", log.status_code == 200
          and any(r["report_id"] == rid and r["kind"] == "run" and r["outcome"] == "OK" and r["user_name"] and r["duration_ms"] is not None for r in body_log.get("runs", []))
          and body_log["summary"]["runs"] >= 3 and body_log["summary"]["median_ms"] is not None and any(p["name"] for p in body_log["people"]), log.content[:300])
    check("the run log narrows by outcome", all(r["outcome"] == "TIMEOUT" for r in world.api("get", "/api/v1/insights/runs/?outcome=TIMEOUT").json()["runs"]))

    # the smaller gaps: Show the SQL, Explain This Report and Change with a Prompt (the provider faked), usage, embed links, linked datasets
    shown_sql = world.api("post", "/api/v1/insights/reports/sql/", draft)
    statements = shown_sql.json().get("statements", []) if shown_sql.status_code == 200 else []
    check("Show the SQL gives the statement with its values written in", shown_sql.status_code == 200 and '"reporting"."rpt_messages"' in statements[0]["sql"] and "%s" not in statements[0]["sql"], shown_sql.content[:300])
    model = {"id": "anthropic:claude-haiku-5-5", "provider": "anthropic", "provider_name": "Claude (Anthropic)", "mark": "A", "model": "claude-haiku-5-5", "label": "Claude Haiku 5.5", "source": "own"}
    offered = {"enabled": True, "reason": None, "models": [model], "default_model": model["id"], "usage": {"used": 0, "limit": None, "words": ""}}
    with mock.patch.object(assistant, "offer", return_value=offered), mock.patch.object(assistant, "effective_providers", return_value={"anthropic": {}}), \
            mock.patch.object(assistant, "call_model", return_value=({"summary": "Fall Launch sent the most.", "points": ["Email carried most of it."], "caution": ""}, {"in": 10, "out": 5})) as asked:
        explained = world.api("post", f"/api/v1/insights/reports/{rid}/explain/", {"model": model["id"], "params": {"period": ["2000-01-01", "2100-12-31"]}})           # every date, not the last 7 days
        sent_rows = asked.call_args.args[4] if asked.call_args else ""
        check("Explain This Report sends the report's rows and answers in a few sentences", explained.status_code == 200 and explained.json()["summary"] == "Fall Launch sent the most."
              and explained.json()["points"] == ["Email carried most of it."] and "Sends by campaign" in sent_rows and "Campaign" in sent_rows, explained.content[:300])
        with mock.patch.object(assistant, "workspace_settings", return_value={"enabled": True, "providers": {}, "explain": False}):
            off = world.api("post", f"/api/v1/insights/reports/{rid}/explain/", {"model": model["id"]})
            check("with the workspace's Explain switch off, explaining is refused in words", off.status_code == 409 and "switched off" in off.json()["detail"])
        asked.return_value = ({"kind": "draft", "definition": {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}]}, "view": {"type": "chart", "chart": "bar"}}, {"in": 10, "out": 5})
        changed_api = world.api("post", "/api/v1/insights/assistant/change/", {**draft, "prompt": "By channel only, as a bar chart", "model": model["id"]})
        check("Change with a Prompt answers the changed definition, checked", changed_api.status_code == 200 and changed_api.json()["kind"] == "draft"
              and [c["field"] for c in changed_api.json()["definition"]["columns"]] == ["channel", "sent"] and changed_api.json()["view"]["chart"] == "bar", changed_api.content[:300])
    prompts = world.sql("SELECT kind, outcome FROM ai_prompt WHERE kind IN ('explain', 'change') ORDER BY created_at", fetch=True)
    check("an explanation and a change are recorded as prompts", ("explain", "DRAFT") in prompts and ("change", "DRAFT") in prompts, str(prompts))
    check("the usage of an opening is recorded once a patch 0061 database has the table", world.api("post", "/api/v1/insights/usage/", {"kind": "report", "id": rid}).status_code == 204)
    usage = world.api("get", "/api/v1/insights/usage/?days=30").json()
    if access.present("insights_usage"):
        check("the usage page counts the opening", any(i["item_id"] == rid and i["opens"] >= 1 for i in usage["items"]), str(usage)[:300])
        check("opening again at once is not counted twice", world.api("post", "/api/v1/insights/usage/", {"kind": "report", "id": rid}).status_code == 204
              and next(i for i in world.api("get", "/api/v1/insights/usage/?days=30").json()["items"] if i["item_id"] == rid)["opens"] == 1)
    else:
        check("before patch 0061 the usage page says it is not ready", usage["ready"] is False)
    if access.present("insights_embed"):
        from django.test import Client
        made = world.api("post", "/api/v1/insights/embeds/", {"kind": "report", "item_id": rid, "label": "Partner site", "days": 30})
        secret_path = made.json().get("path", "") if made.status_code == 201 else ""
        anonymous = Client()
        public = anonymous.get("/api/v1/public" + secret_path + "/") if secret_path else None
        check("an embed link opens the report without signing in, personal data masked", made.status_code == 201 and public is not None and public.status_code == 200
              and public.json()["kind"] == "report" and public.json()["name"] == world.api("get", f"/api/v1/insights/reports/{rid}/").json()["name"], (made.content[:200], public.content[:200] if public is not None else None))
        code = secret_path.split("/")[2] if secret_path else "x"
        check("a wrong secret is 404", anonymous.get(f"/api/v1/public/embed/{code}/wrong-secret/").status_code == 404)
        check("a revoked link is 404", world.api("delete", f"/api/v1/insights/embeds/{made.json()['id']}/").status_code == 204
              and anonymous.get("/api/v1/public" + secret_path + "/").status_code == 404)
    else:
        check("before patch 0061 an embed link cannot be made", world.api("post", "/api/v1/insights/embeds/", {"kind": "report", "item_id": rid}).status_code == 409)
    if links.present():
        joined = world.api("post", "/api/v1/insights/reports/preview/", {"dataset": "messages", "definition": {"columns": [{"field": "campaign_name"}, {"field": "sent", "aggregate": "sum"}],
                                                                                                               "blend": {"dataset": "attributed_orders", "columns": [{"field": "amount", "aggregate": "sum"}]}}})
        check("a report on Messages adds Amount from Attributed Orders", joined.status_code == 200 and [c["key"] for c in joined.json()["columns"]] == ["campaign_name", "sent", "attributed_orders.amount"]
              and joined.json()["linked"]["label"] == "Attributed Orders", joined.content[:300])

    # the visuals: a KPI card's grand total, Key Influencers, the Decomposition Tree, and a dashboard's written summary (the provider faked)
    kpi_run = world.api("post", "/api/v1/insights/reports/preview/", {"dataset": "messages", "definition": {"columns": [{"field": "channel"}, {"field": "sent", "aggregate": "sum"}], "view": {"type": "chart", "chart": "kpi", "target": 10}}})
    check("a KPI card's run carries the engine's grand total", kpi_run.status_code == 200 and any(t["keys"] == {} for t in kpi_run.json().get("totals", [])), kpi_run.content[:300])
    drafted = {"dataset": "messages", "definition": {"columns": [{"field": "campaign_name"}, {"field": "sent", "aggregate": "sum"}]}}
    keyed = world.api("post", "/api/v1/insights/reports/analyze/", {**drafted, "kind": "influencers"})
    body_k = keyed.json() if keyed.status_code == 200 else {}
    check("Key Influencers reads a sum per row and says what raises and lowers it", keyed.status_code == 200 and body_k["kind"] == "influencers" and body_k["measure"]["label"].startswith("Sent Count per")
          and isinstance(body_k["higher"], list) and isinstance(body_k["lower"], list) and "Campaign" in body_k["explained_by"], keyed.content[:300])
    counted = world.api("post", "/api/v1/insights/reports/analyze/", {"dataset": "messages", "definition": {"columns": [{"field": "channel"}, {"field": "messages", "aggregate": "count"}]}, "kind": "influencers"})
    check("a count is refused in words", counted.status_code == 422 and "a count" in counted.json()["detail"], counted.content[:200])
    tree = world.api("post", f"/api/v1/insights/reports/{rid}/analyze/", {"kind": "decompose", "measure": "sent", "path": [], "params": {"period": ["2000-01-01", "2100-12-31"]}})
    body_t = tree.json() if tree.status_code == 200 else {}
    check("the Decomposition Tree starts with the total and the fields to split by", tree.status_code == 200 and body_t["total"] is not None and any(o["field"] == "channel" for o in body_t["explain_by"]), tree.content[:300])
    by_channel = world.api("post", f"/api/v1/insights/reports/{rid}/analyze/", {"kind": "decompose", "measure": "sent", "path": [], "by": "channel", "params": {"period": ["2000-01-01", "2100-12-31"]}})
    rows_c = by_channel.json().get("rows", []) if by_channel.status_code == 200 else []
    check("a split by channel gives the largest first", by_channel.status_code == 200 and rows_c and all((a["measure"] or 0) >= (b["measure"] or 0) for a, b in zip(rows_c, rows_c[1:])), by_channel.content[:300])
    if rows_c:
        under = world.api("post", f"/api/v1/insights/reports/{rid}/analyze/", {"kind": "decompose", "measure": "sent", "path": [{"field": "channel", "value": rows_c[0]["value"]}], "by": "auto_high"})
        check("under a value, High Value picks a field left to split by", under.status_code in (200, 422) and (under.status_code == 422 or (under.json()["chosen"] == "auto_high" and under.json()["by"] != "channel")), under.content[:300])
    check("a split by a field not offered is refused", world.api("post", f"/api/v1/insights/reports/{rid}/analyze/", {"kind": "decompose", "path": [], "by": "nope"}).status_code == 422)
    customised = {"dataset": "messages", "definition": {"columns": [{"field": "c_size"}, {"field": "sent", "aggregate": "sum"}, {"field": "c_email_share", "aggregate": "calc"}],
                                                        "derived": [{"label": "Size", "kind": "rule", "rules": [{"field": "channel", "op": "=", "value": "EMAIL", "label": "Email"}], "other": "Other"},
                                                                    {"label": "Email Share", "kind": "measure", "expression": "DIVIDE(SUM([Sent Count], [Channel] = \"EMAIL\"), SUM([Sent Count])) * 100"}]}}
    keyed_c = world.api("post", "/api/v1/insights/reports/analyze/", {**customised, "kind": "influencers"})
    check("Key Influencers explains by the report's own custom fields", keyed_c.status_code == 200 and "Size" in keyed_c.json()["explained_by"], keyed_c.content[:300])
    formula_c = world.api("post", "/api/v1/insights/reports/analyze/", {**customised, "kind": "influencers", "measure": "c_email_share"})
    check("and explains a measure formula", formula_c.status_code == 200 and formula_c.json()["measure"]["field"] == "c_email_share", formula_c.content[:300])
    split_c = world.api("post", "/api/v1/insights/reports/analyze/", {**customised, "kind": "decompose", "measure": "sent", "path": [], "by": "c_size"})
    check("the Decomposition Tree splits by a custom field", split_c.status_code == 200 and split_c.json()["by"] == "c_size" and split_c.json()["rows"], split_c.content[:300])
    board = world.api("post", "/api/v1/insights/dashboards/", {"name": "Written"}).json()
    put = world.api("put", f"/api/v1/insights/dashboards/{board['id']}/", {"row_version": board["row_version"], "layout": {"rows": [{"width": "1", "tiles": [{"id": "s1", "report_id": rid, "show": "summary"}]}]}})
    check("a summary tile saves", put.status_code == 200, put.content[:200])
    model_s = {"id": "anthropic:claude-haiku-5-5", "provider": "anthropic", "provider_name": "Claude (Anthropic)", "mark": "A", "model": "claude-haiku-5-5", "label": "Claude Haiku 5.5", "source": "own"}
    with mock.patch.object(ai_report, "default_choice", return_value=(model_s, {})), \
            mock.patch.object(assistant, "call_model", return_value=({"summary": "Fall Launch led the month.", "points": ["Email carried it."], "caution": ""}, {"in": 5, "out": 5})) as wrote:
        first = world.api("post", f"/api/v1/insights/dashboards/{board['id']}/tiles/s1/run/", {"period": {"preset": "custom", "from": "2022-01-01", "to": "2030-12-31"}})
        again = world.api("post", f"/api/v1/insights/dashboards/{board['id']}/tiles/s1/run/", {"period": {"preset": "custom", "from": "2022-01-01", "to": "2030-12-31"}})
        calls_s = wrote.call_count
    check("a summary tile writes a few sentences once, then shows the kept ones", first.status_code == 200 and first.json()["kind"] == "summary" and first.json()["summary"] == "Fall Launch led the month."
          and again.json().get("cached") is True and calls_s == 1, (first.content[:300], calls_s))
    board_pdf = world.api("get", f"/api/v1/insights/dashboards/{board['id']}/export.pdf")
    check("the dashboard's PDF carries the kept summary", board_pdf.status_code == 200 and board_pdf.content[:5] == b"%PDF-")
    world.api("delete", f"/api/v1/insights/dashboards/{board['id']}/")

    # dashboards and interaction: a free canvas with a theme, a text tile and a button; PowerPoint and print-ready exports; alerts; comments
    canvas_board = world.api("post", "/api/v1/insights/dashboards/", {"name": "Canvas"}).json()
    laid = world.api("put", f"/api/v1/insights/dashboards/{canvas_board['id']}/", {"row_version": canvas_board["row_version"], "layout": {"rows": [], "canvas": True, "theme": "navy", "cells": [
        {"id": "c1", "report_id": rid, "show": "chart", "x": 0, "y": 0, "w": 8, "h": 4},
        {"id": "c2", "show": "text", "text": "What this board is for.", "title": "Read Me", "x": 8, "y": 0, "w": 4, "h": 2, "hide_on_phone": True},
        {"id": "c3", "show": "button", "label": "Open the Report", "action": {"kind": "report", "id": rid}, "x": 8, "y": 2, "w": 4, "h": 1}]}})
    body_c = laid.json() if laid.status_code == 200 else {}
    check("a dashboard on a free canvas keeps its places, its theme, a text tile and a button", laid.status_code == 200 and body_c["layout"]["canvas"] is True and body_c["layout"]["theme"] == "navy"
          and [c["id"] for c in body_c["layout"]["cells"]] == ["c1", "c2", "c3"] and body_c["layout"]["cells"][1]["hide_on_phone"] is True and body_c["layout"]["cells"][2]["action"]["kind"] == "report", laid.content[:300])
    outside = world.api("put", f"/api/v1/insights/dashboards/{canvas_board['id']}/", {"row_version": body_c.get("row_version", 0), "layout": {"rows": [], "canvas": True, "cells": [
        {"id": "c1", "report_id": rid, "show": "chart", "x": 10, "y": 0, "w": 4, "h": 4}]}})
    check("a tile past the twelve columns is refused where it sits", outside.status_code == 400 and "layout.cells[0]" in json.dumps(outside.json()), outside.content[:200])
    check("a text tile does not run", world.api("post", f"/api/v1/insights/dashboards/{canvas_board['id']}/tiles/c2/run/", {}).status_code == 404)
    deck_d = world.api("get", f"/api/v1/insights/dashboards/{canvas_board['id']}/export.pptx")
    check("the dashboard exports as PowerPoint", deck_d.status_code == 200 and deck_d.content[:2] == b"PK" and deck_d["Content-Type"].endswith("presentationml.presentation"), deck_d.content[:100])
    check("the canvas dashboard's PDF carries the text tile", world.api("get", f"/api/v1/insights/dashboards/{canvas_board['id']}/export.pdf").content[:5] == b"%PDF-")
    world.api("delete", f"/api/v1/insights/dashboards/{canvas_board['id']}/")
    deck_r = world.api("get", f"/api/v1/insights/reports/{rid}/export.pptx")
    check("the report exports as PowerPoint", deck_r.status_code == 200 and deck_r.content[:2] == b"PK" and 'filename="' in deck_r["Content-Disposition"], deck_r.content[:100])
    printed_r = world.api("get", f"/api/v1/insights/reports/{rid}/export.pdf?layout=print&size=letter&orientation=landscape&group=1")
    check("the report prints page by page", printed_r.status_code == 200 and printed_r.content[:5] == b"%PDF-", printed_r.content[:100])
    if alerts.present():
        with mock.patch("apps.platform_ops.mail.is_configured", return_value=True), mock.patch("apps.platform_ops.mail.send", return_value=True) as mailed:
            made_a = world.api("post", "/api/v1/insights/alerts/", {"report_id": rid, "measure": "sent", "op": ">", "threshold": -1, "every": "day", "at_time": "08:00", "name": "Any sends"})
            check("an alert is made on a report's value", made_a.status_code == 201 and made_a.json()["state"] == "UNKNOWN" and made_a.json()["words"].startswith("When "), made_a.content[:300])
            now_a = world.api("post", f"/api/v1/insights/alerts/{made_a.json()['id']}/check/", {})
            again_a = world.api("post", f"/api/v1/insights/alerts/{made_a.json()['id']}/check/", {})
            check("Check Now crosses the threshold and emails once; the next check does not email again", now_a.status_code == 200 and now_a.json()["state"] == "TRIGGERED" and now_a.json()["notified"] is True
                  and again_a.json()["notified"] is False and mailed.call_count == 1, (now_a.content[:300], mailed.call_count))
            listed_a = world.api("get", f"/api/v1/insights/alerts/{made_a.json()['id']}/").json()
            check("the alert keeps its checks", len(listed_a["checks"]) == 2 and listed_a["checks"][0]["met"] is True)
            check("an alert pauses and is deleted", world.api("post", f"/api/v1/insights/alerts/{made_a.json()['id']}/pause/", {}).json()["status"] == "PAUSED"
                  and world.api("delete", f"/api/v1/insights/alerts/{made_a.json()['id']}/").status_code == 204)
        made_c = world.api("post", "/api/v1/insights/comments/", {"kind": "report", "item_id": rid, "body": "September dipped here.", "anchor": {"field": "channel", "value": "SMS", "label": "SMS"}})
        check("a comment on a value of the report", made_c.status_code == 201 and made_c.json()["anchor"] == {"field": "channel", "value": "SMS", "label": "SMS"} and made_c.json()["mine"] is True, made_c.content[:300])
        resolved = world.api("put", f"/api/v1/insights/comments/{made_c.json()['id']}/", {"resolved": True})
        check("a comment is resolved, listed, and deleted by its author", resolved.json()["resolved_at"] is not None
              and any(c["id"] == made_c.json()["id"] for c in world.api("get", f"/api/v1/insights/comments/?kind=report&item={rid}").json()["comments"])
              and world.api("delete", f"/api/v1/insights/comments/{made_c.json()['id']}/").status_code == 204)
    else:
        check("before patch 0062 alerts and comments say so", world.api("post", "/api/v1/insights/comments/", {"kind": "report", "item_id": rid, "body": "x"}).status_code == 409
              and world.api("get", "/api/v1/insights/alerts/").json()["ready"] is False)

    gone = world.api("delete", f"/api/v1/insights/reports/{rid}/")
    check("a deleted report is gone from the list", gone.status_code == 204 and not any(r["id"] == rid for r in world.api("get", "/api/v1/insights/reports/").json()["reports"]))
    other = world.api("get", "/api/v1/insights/reports/00000000-0000-4000-8000-000000000000/")
    check("a report of another workspace, or none, is 404", other.status_code == 404)
    check.done()


def test_connections_office(check):
    """Configuration -> Report Engine Connections: a connection added with its password stored encrypted and never shown,
    Test Connection recording its outcome on the row, and the profile offered on a workspace's Report Engines row."""
    from django.contrib.auth.models import update_last_login
    from django.contrib.auth.signals import user_logged_in
    from django.test import Client
    from django.urls import reverse

    from apps.identity.models import AppUser
    from apps.insights import clickhouse
    from apps.insights.models import ReportEngineConnection
    from apps.platform_ops import crypto

    if not clickhouse.table_present():
        pytest.skip("patch 0054 has not been applied to this database")
    staff = AppUser.objects.using(P).filter(is_staff=True, is_active=True, is_platform_admin=True).order_by("created_at").first() \
        or AppUser.objects.using(P).filter(is_staff=True, is_active=True).order_by("created_at").first()
    user_logged_in.disconnect(update_last_login, dispatch_uid="update_last_login")
    try:
        office = Client()
        office.force_login(staff)
        page = office.get("/admin/insights/reportengineconnection/")
        check("the Connections page opens", page.status_code == 200, str(page.status_code))
        form = {"profile": "analytics-test", "engine": "clickhouse", "url": "127.0.0.1:9", "username": "ch_user", "password": "s3cret", "database_name": "cep_reporting",
                "status": "ACTIVE", "notes": "", "_save": "Save"}
        refused = office.post("/admin/insights/reportengineconnection/add/", form)
        check("a URL without http is refused with the reason", refused.status_code == 200 and b"HTTP interface" in refused.content, str(refused.status_code))
        added = office.post("/admin/insights/reportengineconnection/add/", {**form, "url": "http://127.0.0.1:9"})
        row = ReportEngineConnection.objects.using(P).filter(profile="analytics-test").first()
        check("a connection is added with its password encrypted", added.status_code == 302 and row is not None and crypto.decrypt(row.password_enc) == "s3cret", str(added.status_code))
        named = clickhouse.connection_named("analytics-test")
        check("the profile resolves to a connection with the password decrypted", named is not None and named.password == "s3cret" and named.url == "http://127.0.0.1:9")
        change = office.get(f"/admin/insights/reportengineconnection/{row.pk}/change/")
        check("the row's page never shows the password", change.status_code == 200 and b"s3cret" not in change.content and b"Test Connection" in change.content, str(change.status_code))
        kept = office.post(f"/admin/insights/reportengineconnection/{row.pk}/change/", {**form, "url": "http://127.0.0.1:9", "password": "", "notes": "the EU server"})
        row.refresh_from_db(using=P)
        check("saving without a password keeps the stored one", kept.status_code == 302 and row.notes == "the EU server" and crypto.decrypt(row.password_enc) == "s3cret", str(kept.status_code))
        tested = office.post(reverse("admin:insights_reportengineconnection_test_connection", args=[row.pk]))
        row.refresh_from_db(using=P)
        check("Test Connection records the outcome on the row", tested.status_code == 302 and row.last_test_ok is False and row.last_tested_at is not None and row.last_test_error, str(tested.status_code))
        if _tables_present():
            with connections[P].cursor() as cursor:
                cursor.execute("SELECT tenant_id FROM report_engine_setting LIMIT 1")
                found = cursor.fetchone()
            if found:
                engines = office.get(f"/admin/insights/reportenginesetting/{found[0]}/change/")
                check("a workspace's Report Engines row offers the profile and the environment", engines.status_code == 200 and b"analytics-test" in engines.content
                      and b"The server&#x27;s environment" in engines.content or b"The server's environment" in engines.content, str(engines.status_code))
    finally:
        user_logged_in.connect(update_last_login, dispatch_uid="update_last_login")
    check.done()
